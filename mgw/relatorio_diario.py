from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import (
    Alignment,
    Border,
    Font,
    PatternFill,
    Side,
)
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import (
    Table,
    TableStyleInfo,
)

from mgw.dashboard_inteligente import DashboardInteligente
from mgw.relatorio_excel import RelatorioExcel


class RelatorioDiario:
    def __init__(
        self,
        banco,
        pasta_saida="relatorios/diarios",
    ):
        self.banco = banco
        self.pasta_saida = Path(pasta_saida)

    def gerar(
        self,
        naps,
        nomes_servidores,
        inicio_periodo,
        fim_periodo,
    ):
        """
        Gera o relatório normal e acrescenta:

        - HISTÓRICO ROTAS
        - HISTÓRICO SERVIDORES
        """
        gerador_base = RelatorioExcel(
            pasta_saida=self.pasta_saida
        )

        caminho = gerador_base.gerar(
            naps=naps,
            nomes_servidores=nomes_servidores,
        )

        workbook = load_workbook(caminho)

        todos_eventos_rotas = (
            self.banco.listar_eventos(
                limite=100000
            )
        )

        todos_eventos_servidores = (
            self.banco.listar_eventos_servidores(
                limite=100000
            )
        )

        eventos_rotas = self._eventos_do_periodo(
            eventos=todos_eventos_rotas,
            inicio_periodo=inicio_periodo,
            fim_periodo=fim_periodo,
        )

        eventos_servidores = (
            self._eventos_do_periodo(
                eventos=todos_eventos_servidores,
                inicio_periodo=inicio_periodo,
                fim_periodo=fim_periodo,
            )
        )

        self._criar_historico_rotas(
            workbook=workbook,
            eventos=eventos_rotas,
            fim_periodo=fim_periodo,
        )

        self._criar_historico_servidores(
            workbook=workbook,
            eventos=eventos_servidores,
            fim_periodo=fim_periodo,
        )

        dashboard = DashboardInteligente()
        dados_dashboard = dashboard.calcular(
            naps=naps,
            nomes_servidores=nomes_servidores,
            eventos_rotas=todos_eventos_rotas,
            eventos_servidores=(
                todos_eventos_servidores
            ),
            inicio_periodo=inicio_periodo,
            fim_periodo=fim_periodo,
            banco=self.banco,
        )

        dashboard.criar_abas(
            workbook=workbook,
            dados=dados_dashboard,
        )

        workbook.save(caminho)

        return caminho

    @staticmethod
    def _eventos_do_periodo(
        eventos,
        inicio_periodo,
        fim_periodo,
    ):
        """
        Mantém eventos que tiveram alguma sobreposição
        com o período analisado.

        Um evento iniciado antes da meia-noite, mas que
        permaneceu aberto durante o dia, também aparece.
        """
        resultado = []

        for evento in eventos:
            caiu_em = RelatorioDiario._converter_data(
                evento.get("caiu_em")
            )

            voltou_em = (
                RelatorioDiario._converter_data(
                    evento.get("voltou_em")
                )
            )

            if caiu_em is None:
                continue

            if caiu_em > fim_periodo:
                continue

            if (
                voltou_em is not None
                and voltou_em < inicio_periodo
            ):
                continue

            resultado.append(evento)

        return resultado

    def _criar_historico_rotas(
        self,
        workbook,
        eventos,
        fim_periodo,
    ):
        nome_aba = "HISTÓRICO ROTAS"

        if nome_aba in workbook.sheetnames:
            del workbook[nome_aba]

        planilha = workbook.create_sheet(
            title=nome_aba
        )

        planilha.sheet_view.showGridLines = False

        cabecalhos = [
            "Servidor",
            "Rota",
            "Endereço",
            "Caiu",
            "Voltou",
            "Tempo parado",
            "Situação",
        ]

        planilha.append(cabecalhos)

        eventos_ordenados = sorted(
            eventos,
            key=lambda item: (
                item.get("caiu_em") or ""
            ),
            reverse=True,
        )

        for evento in eventos_ordenados:
            caiu_em = self._converter_data(
                evento.get("caiu_em")
            )

            if caiu_em is None:
                continue

            voltou_em = self._converter_data(
                evento.get("voltou_em")
            )

            fim_evento = voltou_em or fim_periodo

            duracao = max(
                0,
                int(
                    (
                        fim_evento - caiu_em
                    ).total_seconds()
                ),
            )

            situacao = (
                "Em andamento"
                if voltou_em is None
                else "Recuperado"
            )

            planilha.append(
                [
                    evento.get("servidor"),
                    evento.get("rota"),
                    evento.get("endereco"),
                    caiu_em,
                    voltou_em,
                    self._formatar_duracao(
                        duracao
                    ),
                    situacao,
                ]
            )

        self._formatar_historico(
            planilha=planilha,
            nome_tabela="TabelaHistoricoRotas",
            quantidade_colunas=7,
            colunas_data=(4, 5),
            larguras={
                "A": 18,
                "B": 55,
                "C": 30,
                "D": 22,
                "E": 22,
                "F": 22,
                "G": 20,
            },
        )

    def _criar_historico_servidores(
        self,
        workbook,
        eventos,
        fim_periodo,
    ):
        nome_aba = "HISTÓRICO SERVIDORES"

        if nome_aba in workbook.sheetnames:
            del workbook[nome_aba]

        planilha = workbook.create_sheet(
            title=nome_aba
        )

        planilha.sheet_view.showGridLines = False

        cabecalhos = [
            "Servidor",
            "Endereço",
            "Ficou inacessível",
            "Restabelecido",
            "Tempo parado",
            "Falhas consecutivas",
            "Situação",
        ]

        planilha.append(cabecalhos)

        eventos_ordenados = sorted(
            eventos,
            key=lambda item: (
                item.get("caiu_em") or ""
            ),
            reverse=True,
        )

        for evento in eventos_ordenados:
            caiu_em = self._converter_data(
                evento.get("caiu_em")
            )

            if caiu_em is None:
                continue

            voltou_em = self._converter_data(
                evento.get("voltou_em")
            )

            fim_evento = voltou_em or fim_periodo

            duracao = max(
                0,
                int(
                    (
                        fim_evento - caiu_em
                    ).total_seconds()
                ),
            )

            situacao = (
                "Inacessível"
                if voltou_em is None
                else "Restabelecido"
            )

            planilha.append(
                [
                    evento.get("servidor"),
                    evento.get("endereco"),
                    caiu_em,
                    voltou_em,
                    self._formatar_duracao(
                        duracao
                    ),
                    evento.get(
                        "falhas_consecutivas"
                    ),
                    situacao,
                ]
            )

        self._formatar_historico(
            planilha=planilha,
            nome_tabela=(
                "TabelaHistoricoServidores"
            ),
            quantidade_colunas=7,
            colunas_data=(3, 4),
            larguras={
                "A": 18,
                "B": 28,
                "C": 22,
                "D": 22,
                "E": 22,
                "F": 22,
                "G": 20,
            },
        )

    @staticmethod
    def _formatar_historico(
        planilha,
        nome_tabela,
        quantidade_colunas,
        colunas_data,
        larguras,
    ):
        preenchimento_cabecalho = PatternFill(
            fill_type="solid",
            fgColor="37474F",
        )

        fonte_cabecalho = Font(
            color="FFFFFF",
            bold=True,
        )

        preenchimento_normal = PatternFill(
            fill_type="solid",
            fgColor="E8F5E9",
        )

        preenchimento_aberto = PatternFill(
            fill_type="solid",
            fgColor="F4CCCC",
        )

        borda_fina = Side(
            style="thin",
            color="D9D9D9",
        )

        borda = Border(
            left=borda_fina,
            right=borda_fina,
            top=borda_fina,
            bottom=borda_fina,
        )

        for celula in planilha[1]:
            celula.fill = preenchimento_cabecalho
            celula.font = fonte_cabecalho
            celula.border = borda
            celula.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True,
            )

        planilha.row_dimensions[1].height = 30

        for linha in planilha.iter_rows(
            min_row=2,
        ):
            situacao = str(
                linha[
                    quantidade_colunas - 1
                ].value
                or ""
            )

            evento_aberto = situacao in {
                "Em andamento",
                "Inacessível",
            }

            preenchimento = (
                preenchimento_aberto
                if evento_aberto
                else preenchimento_normal
            )

            for celula in linha:
                celula.fill = preenchimento
                celula.border = borda
                celula.alignment = Alignment(
                    vertical="center",
                    wrap_text=True,
                )

            planilha.row_dimensions[
                linha[0].row
            ].height = 34

        planilha.freeze_panes = "A2"
        planilha.auto_filter.ref = (
            planilha.dimensions
        )

        for coluna, largura in larguras.items():
            planilha.column_dimensions[
                coluna
            ].width = largura

        for numero_linha in range(
            2,
            planilha.max_row + 1,
        ):
            for numero_coluna in colunas_data:
                celula = planilha.cell(
                    row=numero_linha,
                    column=numero_coluna,
                )

                if isinstance(
                    celula.value,
                    datetime,
                ):
                    celula.number_format = (
                        "dd/mm/yyyy hh:mm:ss"
                    )

                    celula.alignment = Alignment(
                        horizontal="center",
                        vertical="center",
                    )

        # Centraliza colunas curtas, mantendo rota
        # e endereço alinhados à esquerda.
        for numero_linha in range(
            2,
            planilha.max_row + 1,
        ):
            for numero_coluna in range(
                1,
                quantidade_colunas + 1,
            ):
                if numero_coluna in colunas_data:
                    continue

                celula = planilha.cell(
                    row=numero_linha,
                    column=numero_coluna,
                )

                if numero_coluna in {
                    1,
                    5,
                    6,
                    7,
                }:
                    celula.alignment = Alignment(
                        horizontal="center",
                        vertical="center",
                        wrap_text=True,
                    )

        if planilha.max_row >= 2:
            ultima_coluna = get_column_letter(
                quantidade_colunas
            )

            tabela = Table(
                displayName=nome_tabela,
                ref=(
                    f"A1:{ultima_coluna}"
                    f"{planilha.max_row}"
                ),
            )

            estilo = TableStyleInfo(
                name="TableStyleMedium2",
                showFirstColumn=False,
                showLastColumn=False,
                showRowStripes=False,
                showColumnStripes=False,
            )

            tabela.tableStyleInfo = estilo
            planilha.add_table(tabela)

    @staticmethod
    def _converter_data(valor):
        if valor is None:
            return None

        if isinstance(valor, datetime):
            return valor

        try:
            return datetime.fromisoformat(
                str(valor)
            )

        except (
            ValueError,
            TypeError,
        ):
            return None

    @staticmethod
    def _formatar_duracao(segundos):
        segundos = max(
            0,
            int(segundos or 0),
        )

        dias, resto = divmod(
            segundos,
            86400,
        )

        horas, resto = divmod(
            resto,
            3600,
        )

        minutos, segundos = divmod(
            resto,
            60,
        )

        partes = []

        if dias:
            partes.append(
                f"{dias} dia"
                if dias == 1
                else f"{dias} dias"
            )

        if horas:
            partes.append(
                f"{horas} h"
            )

        if minutos:
            partes.append(
                f"{minutos} min"
            )

        if segundos or not partes:
            partes.append(
                f"{segundos} s"
            )

        return " ".join(partes)