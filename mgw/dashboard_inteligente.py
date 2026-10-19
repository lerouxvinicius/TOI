from collections import defaultdict
from datetime import datetime, timedelta

from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


class DashboardInteligente:
    """
    Consolida os indicadores do MGW Monitor.

    Regras principais:
    - Disponibilidade dos MGWs: período do dia do relatório.
    - Indisponibilidade de MGW aberto: duração real desde o início do evento.
    - Ranking principal: IP Peer em janela móvel de 7 dias.
    - Um incidente do IP é contado uma única vez, mesmo quando várias
      rotas ou MGWs do mesmo endereço ficam indisponíveis juntos.
    """

    JANELA_RANKING_DIAS = 7

    def calcular(
        self,
        naps,
        nomes_servidores,
        eventos_rotas,
        eventos_servidores,
        inicio_periodo,
        fim_periodo,
        banco=None,
    ):
        eventos_rotas = self._complementar_eventos_atuais(
            eventos=list(eventos_rotas),
            naps=naps,
            banco=banco,
            fim_periodo=fim_periodo,
        )

        periodo_segundos = max(
            1,
            int((fim_periodo - inicio_periodo).total_seconds()),
        )

        inicio_ranking = fim_periodo - timedelta(
            days=self.JANELA_RANKING_DIAS
        )

        rotas_por_servidor = defaultdict(set)
        rotas_por_endereco = defaultdict(set)
        servidores_por_endereco = defaultdict(set)

        for nap in naps:
            servidor = str(
                nap.servidor or "Não informado"
            )
            rota = str(
                nap.nome or "Rota não informada"
            )
            endereco = str(
                nap.endereco() or "Não informado"
            )

            rotas_por_servidor[servidor].add(rota)
            rotas_por_endereco[endereco].add(rota)
            servidores_por_endereco[endereco].add(
                servidor
            )

        eventos_rotas_dia = self._eventos_sobrepostos(
            eventos_rotas,
            inicio_periodo,
            fim_periodo,
        )

        eventos_servidores_dia = (
            self._eventos_sobrepostos(
                eventos_servidores,
                inicio_periodo,
                fim_periodo,
            )
        )

        eventos_rotas_7_dias = (
            self._eventos_sobrepostos(
                eventos_rotas,
                inicio_ranking,
                fim_periodo,
            )
        )

        servidores = self._estatisticas_servidores(
            nomes_servidores=nomes_servidores,
            rotas_por_servidor=rotas_por_servidor,
            eventos_dia=eventos_servidores_dia,
            todos_eventos=eventos_servidores,
            inicio_periodo=inicio_periodo,
            fim_periodo=fim_periodo,
            periodo_segundos=periodo_segundos,
        )

        rotas_7_dias = self._estatisticas_rotas(
            eventos=eventos_rotas_7_dias,
            inicio_periodo=inicio_ranking,
            fim_periodo=fim_periodo,
        )

        enderecos_inativos_atuais = {
            str(
                nap.endereco() or "Não informado"
            ).strip().casefold()
            for nap in naps
            if nap.esta_inativo()
        }

        enderecos_7_dias = self._estatisticas_enderecos(
            eventos=eventos_rotas_7_dias,
            rotas_por_endereco=rotas_por_endereco,
            servidores_por_endereco=(
                servidores_por_endereco
            ),
            inicio_periodo=inicio_ranking,
            fim_periodo=fim_periodo,
            enderecos_inativos_atuais=(
                enderecos_inativos_atuais
            ),
        )

        ranking_rotas = sorted(
            rotas_7_dias,
            key=lambda item: (
                -item["quedas"],
                -item["tempo_total_segundos"],
                item["rota"].casefold(),
            ),
        )

        ranking_enderecos = sorted(
            enderecos_7_dias,
            key=lambda item: (
                -item["quedas"],
                -item["tempo_total_segundos"],
                item["endereco"].casefold(),
            ),
        )

        ranking_servidores = sorted(
            servidores,
            key=lambda item: (
                item["disponibilidade"],
                -item["quedas"],
                item["servidor"].casefold(),
            ),
        )

        maior_indisponibilidade = (
            self._maior_indisponibilidade_dia(
                eventos=eventos_rotas_dia,
                inicio_periodo=inicio_periodo,
                fim_periodo=fim_periodo,
            )
        )

        disponibilidade_geral = (
            sum(
                item["disponibilidade"]
                for item in servidores
            )
            / len(servidores)
            if servidores
            else 100.0
        )

        total_quedas_rotas = sum(
            1
            for evento in eventos_rotas_dia
            if self._esta_no_periodo(
                evento.get("caiu_em"),
                inicio_periodo,
                fim_periodo,
            )
        )

        total_recuperacoes_rotas = sum(
            1
            for evento in eventos_rotas_dia
            if self._esta_no_periodo(
                evento.get("voltou_em"),
                inicio_periodo,
                fim_periodo,
            )
        )

        return {
            "inicio_periodo": inicio_periodo,
            "fim_periodo": fim_periodo,
            "inicio_ranking": inicio_ranking,
            "fim_ranking": fim_periodo,
            "periodo_segundos": periodo_segundos,
            "disponibilidade_geral": (
                disponibilidade_geral
            ),
            "total_quedas_rotas": total_quedas_rotas,
            "total_recuperacoes_rotas": (
                total_recuperacoes_rotas
            ),
            "servidores": servidores,
            "rotas": rotas_7_dias,
            "enderecos": enderecos_7_dias,
            "ranking_rotas": ranking_rotas,
            "ranking_enderecos": ranking_enderecos,
            "ranking_servidores": ranking_servidores,
            "maior_indisponibilidade": (
                maior_indisponibilidade
            ),
        }

    def criar_abas(self, workbook, dados):
        for nome in (
            "DASHBOARD",
            "ESTATÍSTICAS",
            "ENDEREÇOS",
            "ROTAS",
            "SERVIDORES",
        ):
            if nome in workbook.sheetnames:
                del workbook[nome]

        self._criar_dashboard(workbook, dados)
        self._criar_estatisticas(workbook, dados)
        self._criar_enderecos(workbook, dados)
        self._criar_rotas(workbook, dados)
        self._criar_servidores(workbook, dados)

        workbook._sheets.insert(
            0,
            workbook._sheets.pop(
                workbook.sheetnames.index("DASHBOARD")
            ),
        )

    def montar_mensagem_telegram(
        self,
        telegram,
        dados,
        limite_ranking=5,
    ):
        servidores = sorted(
            dados["servidores"],
            key=lambda item: (
                item["servidor"].casefold()
            ),
        )

        linhas = [
            "📈 <b>MGW MONITOR — DASHBOARD DO DIA</b>",
            "",
            "🖥 <b>DISPONIBILIDADE DOS MGWs</b>",
            "",
        ]

        for item in servidores:
            linhas.append(
                f"{telegram.escapar(item['servidor'])}   "
                f"{self._formatar_percentual(item['disponibilidade'])}"
            )

        linhas.extend(
            [
                "",
                "⏱ <b>INDISPONIBILIDADE DOS MGWs</b>",
                "",
            ]
        )

        for item in servidores:
            if item["indisponivel_atual"]:
                tempo = self._formatar_duracao_horas(
                    item[
                        "indisponibilidade_atual_segundos"
                    ]
                )
            else:
                tempo = self._formatar_duracao(
                    item["tempo_medio_segundos"]
                )

            linhas.append(
                f"{telegram.escapar(item['servidor'])}   "
                f"{tempo}"
            )

        linhas.extend(
            [
                "",
                "🌐 <b>TOP IP PEER — 7 DIAS</b>",
                "",
            ]
        )

        ranking_com_quedas = [
            item
            for item in dados["ranking_enderecos"]
            if item["quedas"] > 0
        ]

        ranking_enderecos = (
            ranking_com_quedas[:limite_ranking]
        )

        if ranking_enderecos:
            for posicao, item in enumerate(
                ranking_enderecos,
                start=1,
            ):
                quantidade = (
                    "queda"
                    if item["quedas"] == 1
                    else "quedas"
                )

                situacao = (
                    " - ainda indisponível"
                    if item.get("indisponivel_atual")
                    else ""
                )

                linhas.extend(
                    [
                        self._medalha(posicao)
                        + " "
                        + telegram.escapar(
                            item["endereco"]
                        ),
                        (
                            f"{item['quedas']} {quantidade}"
                            f"{situacao}"
                        ),
                        "",
                    ]
                )

            if len(ranking_com_quedas) > limite_ranking:
                linhas.extend(
                    [
                        (
                            "📎 Demais IP Peers com ocorrências "
                            "estão disponíveis no Excel."
                        ),
                        "",
                    ]
                )
        else:
            linhas.extend(
                ["Nenhuma queda registrada.", ""]
            )

        return "\n".join(linhas).strip()

    def _estatisticas_servidores(
        self,
        nomes_servidores,
        rotas_por_servidor,
        eventos_dia,
        todos_eventos,
        inicio_periodo,
        fim_periodo,
        periodo_segundos,
    ):
        eventos_dia_por_servidor = defaultdict(list)
        todos_por_servidor = defaultdict(list)

        for evento in eventos_dia:
            servidor = str(
                evento.get("servidor")
                or "Não informado"
            )
            eventos_dia_por_servidor[servidor].append(
                evento
            )

        for evento in todos_eventos:
            servidor = str(
                evento.get("servidor")
                or "Não informado"
            )
            todos_por_servidor[servidor].append(
                evento
            )

        resultado = []

        for servidor in nomes_servidores:
            lista_dia = eventos_dia_por_servidor.get(
                servidor,
                [],
            )

            lista_todos = todos_por_servidor.get(
                servidor,
                [],
            )

            quedas = sum(
                1
                for evento in lista_dia
                if self._esta_no_periodo(
                    evento.get("caiu_em"),
                    inicio_periodo,
                    fim_periodo,
                )
            )

            tempo_total_dia = self._duracao_uniao(
                lista_dia,
                inicio_periodo,
                fim_periodo,
            )

            disponibilidade = max(
                0.0,
                100.0
                * (
                    1
                    - tempo_total_dia
                    / periodo_segundos
                ),
            )

            eventos_encerrados_dia = [
                evento
                for evento in lista_dia
                if self._converter_data(
                    evento.get("voltou_em")
                )
                is not None
            ]

            duracoes_encerradas = [
                self._duracao_evento_completa(evento)
                for evento in eventos_encerrados_dia
            ]

            duracoes_encerradas = [
                valor
                for valor in duracoes_encerradas
                if valor is not None
            ]

            tempo_medio = (
                int(
                    sum(duracoes_encerradas)
                    / len(duracoes_encerradas)
                )
                if duracoes_encerradas
                else 0
            )

            evento_aberto = self._evento_aberto_mais_antigo(
                lista_todos
            )

            indisponivel_atual = (
                evento_aberto is not None
            )

            indisponibilidade_atual = 0
            indisponivel_desde = None

            if evento_aberto is not None:
                indisponivel_desde = self._converter_data(
                    evento_aberto.get("caiu_em")
                )

                if indisponivel_desde is not None:
                    indisponibilidade_atual = max(
                        0,
                        int(
                            (
                                fim_periodo
                                - indisponivel_desde
                            ).total_seconds()
                        ),
                    )

            resultado.append(
                {
                    "servidor": servidor,
                    "rotas_monitoradas": len(
                        rotas_por_servidor.get(
                            servidor,
                            set(),
                        )
                    ),
                    "quedas": quedas,
                    "tempo_total_segundos": (
                        tempo_total_dia
                    ),
                    "tempo_medio_segundos": tempo_medio,
                    "disponibilidade": disponibilidade,
                    "indisponivel_atual": (
                        indisponivel_atual
                    ),
                    "indisponivel_desde": (
                        indisponivel_desde
                    ),
                    "indisponibilidade_atual_segundos": (
                        indisponibilidade_atual
                    ),
                }
            )

        return resultado

    def _estatisticas_rotas(
        self,
        eventos,
        inicio_periodo,
        fim_periodo,
    ):
        agrupados = defaultdict(list)
        servidores = defaultdict(set)
        enderecos = defaultdict(set)

        for evento in eventos:
            rota = str(
                evento.get("rota")
                or "Rota não informada"
            )
            agrupados[rota].append(evento)
            servidores[rota].add(
                str(
                    evento.get("servidor")
                    or "Não informado"
                )
            )
            enderecos[rota].add(
                str(
                    evento.get("endereco")
                    or "Não informado"
                )
            )

        periodo_segundos = max(
            1,
            int((fim_periodo - inicio_periodo).total_seconds()),
        )

        resultado = []

        for rota, lista in agrupados.items():
            quedas = sum(
                1
                for evento in lista
                if self._esta_no_periodo(
                    evento.get("caiu_em"),
                    inicio_periodo,
                    fim_periodo,
                )
            )

            duracoes = [
                self._duracao_sobreposta(
                    evento,
                    inicio_periodo,
                    fim_periodo,
                )
                for evento in lista
            ]

            tempo_total = self._duracao_uniao(
                lista,
                inicio_periodo,
                fim_periodo,
            )

            media = (
                int(sum(duracoes) / len(duracoes))
                if duracoes
                else 0
            )

            disponibilidade = max(
                0.0,
                100.0
                * (
                    1
                    - tempo_total
                    / periodo_segundos
                ),
            )

            resultado.append(
                {
                    "rota": rota,
                    "quedas": quedas,
                    "enderecos": sorted(
                        enderecos[rota],
                        key=str.casefold,
                    ),
                    "servidores": sorted(
                        servidores[rota],
                        key=str.casefold,
                    ),
                    "tempo_total_segundos": tempo_total,
                    "tempo_medio_segundos": media,
                    "disponibilidade": disponibilidade,
                }
            )

        return resultado

    def _estatisticas_enderecos(
        self,
        eventos,
        rotas_por_endereco,
        servidores_por_endereco,
        inicio_periodo,
        fim_periodo,
        enderecos_inativos_atuais=None,
    ):
        enderecos_inativos_atuais = (
            enderecos_inativos_atuais or set()
        )
        agrupados = defaultdict(list)

        for evento in eventos:
            endereco = str(
                evento.get("endereco")
                or "Não informado"
            )

            agrupados[endereco].append(evento)

            rotas_por_endereco[endereco].add(
                str(
                    evento.get("rota")
                    or "Rota não informada"
                )
            )

            servidores_por_endereco[endereco].add(
                str(
                    evento.get("servidor")
                    or "Não informado"
                )
            )

        todos_enderecos = (
            set(rotas_por_endereco)
            | set(agrupados)
        )

        periodo_segundos = max(
            1,
            int((fim_periodo - inicio_periodo).total_seconds()),
        )

        resultado = []

        for endereco in todos_enderecos:
            lista = agrupados.get(endereco, [])

            # Um mesmo IP pode sustentar várias rotas e estar
            # configurado em vários MGWs. Para o ranking do IP Peer,
            # eventos simultâneos/sobrepostos contam como um único
            # incidente do endereço.
            quedas = self._quantidade_incidentes_uniao(
                lista,
                inicio_periodo,
                fim_periodo,
            )

            tempo_total = self._duracao_uniao(
                lista,
                inicio_periodo,
                fim_periodo,
            )

            media = (
                int(tempo_total / quedas)
                if quedas > 0
                else 0
            )

            disponibilidade = max(
                0.0,
                100.0
                * (
                    1
                    - tempo_total
                    / periodo_segundos
                ),
            )

            resultado.append(
                {
                    "endereco": endereco,
                    "quedas": quedas,
                    "rotas": sorted(
                        rotas_por_endereco.get(
                            endereco,
                            set(),
                        ),
                        key=str.casefold,
                    ),
                    "servidores": sorted(
                        servidores_por_endereco.get(
                            endereco,
                            set(),
                        ),
                        key=str.casefold,
                    ),
                    "tempo_total_segundos": tempo_total,
                    "tempo_medio_segundos": media,
                    "disponibilidade": disponibilidade,
                    "indisponivel_atual": (
                        endereco.strip().casefold()
                        in enderecos_inativos_atuais
                    ),
                }
            )

        return resultado

    def _maior_indisponibilidade_dia(
        self,
        eventos,
        inicio_periodo,
        fim_periodo,
    ):
        maior = None
        maior_duracao = -1

        for evento in eventos:
            duracao = self._duracao_sobreposta(
                evento,
                inicio_periodo,
                fim_periodo,
            )

            if duracao > maior_duracao:
                maior_duracao = duracao
                maior = {
                    "servidor": (
                        evento.get("servidor")
                        or "Não informado"
                    ),
                    "rota": (
                        evento.get("rota")
                        or "Rota não informada"
                    ),
                    "endereco": (
                        evento.get("endereco")
                        or "Não informado"
                    ),
                    "duracao_segundos": duracao,
                }

        return maior

    def _criar_dashboard(self, workbook, dados):
        ws = workbook.create_sheet("DASHBOARD")
        ws.sheet_view.showGridLines = False

        ws.merge_cells("A1:H1")
        ws["A1"] = "MGW Monitor — Dashboard Inteligente"
        ws["A1"].font = Font(
            size=20,
            bold=True,
            color="FFFFFF",
        )
        ws["A1"].fill = PatternFill(
            "solid",
            fgColor="1B5E20",
        )
        ws["A1"].alignment = Alignment(
            horizontal="center",
            vertical="center",
        )
        ws.row_dimensions[1].height = 32

        ws.merge_cells("A2:H2")
        ws["A2"] = (
            "Período diário: "
            f"{dados['inicio_periodo']:%d/%m/%Y %H:%M:%S} até "
            f"{dados['fim_periodo']:%d/%m/%Y %H:%M:%S} | "
            "Rankings: últimos 7 dias"
        )
        ws["A2"].alignment = Alignment(
            horizontal="center"
        )
        ws["A2"].font = Font(italic=True)

        indicadores = [
            (
                "Quedas de rotas no dia",
                dados["total_quedas_rotas"],
            ),
            (
                "Recuperações no dia",
                dados["total_recuperacoes_rotas"],
            ),
            (
                "MGWs monitorados",
                len(dados["servidores"]),
            ),
        ]

        for indice, (titulo, valor) in enumerate(
            indicadores,
            start=0,
        ):
            col_ini = 1 + indice * 2
            ws.merge_cells(
                start_row=4,
                start_column=col_ini,
                end_row=4,
                end_column=col_ini + 1,
            )
            ws.merge_cells(
                start_row=5,
                start_column=col_ini,
                end_row=5,
                end_column=col_ini + 1,
            )

            c1 = ws.cell(4, col_ini, titulo)
            c2 = ws.cell(5, col_ini, valor)

            for celula in (c1, c2):
                celula.alignment = Alignment(
                    horizontal="center",
                    vertical="center",
                )

            c1.font = Font(
                bold=True,
                color="FFFFFF",
            )
            c1.fill = PatternFill(
                "solid",
                fgColor="2E7D32",
            )
            c2.font = Font(
                size=16,
                bold=True,
            )
            c2.fill = PatternFill(
                "solid",
                fgColor="E8F5E9",
            )

        linha = 8
        ws.cell(
            linha,
            1,
            "Disponibilidade por MGW",
        )
        ws.cell(linha, 1).font = Font(
            size=14,
            bold=True,
            color="FFFFFF",
        )
        ws.cell(linha, 1).fill = PatternFill(
            "solid",
            fgColor="1565C0",
        )
        ws.merge_cells(
            start_row=linha,
            start_column=1,
            end_row=linha,
            end_column=5,
        )

        linha += 1
        cabecalhos = [
            "Servidor",
            "Disponibilidade",
            "Tempo indisponível no dia",
            "Indisponibilidade atual / média",
            "Quedas no dia",
        ]

        for col, valor in enumerate(
            cabecalhos,
            1,
        ):
            ws.cell(linha, col, valor)

        self._formatar_cabecalho(
            ws,
            linha,
            1,
            5,
        )

        linha_inicio_servidores = linha + 1

        for item in sorted(
            dados["servidores"],
            key=lambda x: x[
                "servidor"
            ].casefold(),
        ):
            linha += 1
            ws.cell(linha, 1, item["servidor"])
            ws.cell(
                linha,
                2,
                item["disponibilidade"] / 100,
            )
            ws.cell(linha, 2).number_format = "0.00%"
            ws.cell(
                linha,
                3,
                self._formatar_duracao(
                    item["tempo_total_segundos"]
                ),
            )

            if item["indisponivel_atual"]:
                tempo_referencia = (
                    self._formatar_duracao_horas(
                        item[
                            "indisponibilidade_atual_segundos"
                        ]
                    )
                )
            else:
                tempo_referencia = self._formatar_duracao(
                    item["tempo_medio_segundos"]
                )

            ws.cell(
                linha,
                4,
                tempo_referencia,
            )
            ws.cell(
                linha,
                5,
                item["quedas"],
            )

        linha_fim_servidores = linha
        self._formatar_corpo(
            ws,
            linha_inicio_servidores,
            linha_fim_servidores,
            1,
            5,
        )

        linha += 3
        ws.cell(
            linha,
            1,
            "Top 10 IP Peer — 7 dias",
        )
        ws.cell(linha, 1).font = Font(
            size=14,
            bold=True,
            color="FFFFFF",
        )
        ws.cell(linha, 1).fill = PatternFill(
            "solid",
            fgColor="EF6C00",
        )
        ws.merge_cells(
            start_row=linha,
            start_column=1,
            end_row=linha,
            end_column=4,
        )

        linha += 1
        for col, valor in enumerate(
            [
                "Posição",
                "IP Peer",
                "Quedas",
                "Tempo total (7 dias)",
            ],
            1,
        ):
            ws.cell(linha, col, valor)

        self._formatar_cabecalho(
            ws,
            linha,
            1,
            4,
        )

        inicio_enderecos = linha + 1

        for posicao, item in enumerate(
            dados["ranking_enderecos"][:10],
            start=1,
        ):
            linha += 1
            ws.cell(linha, 1, posicao)
            ws.cell(linha, 2, item["endereco"])
            ws.cell(linha, 3, item["quedas"])
            ws.cell(
                linha,
                4,
                self._formatar_duracao(
                    item["tempo_total_segundos"]
                ),
            )

        self._formatar_corpo(
            ws,
            inicio_enderecos,
            linha,
            1,
            4,
        )

        larguras = {
            "A": 18,
            "B": 55,
            "C": 25,
            "D": 30,
            "E": 15,
            "F": 18,
            "G": 18,
            "H": 18,
        }

        for coluna, largura in larguras.items():
            ws.column_dimensions[coluna].width = largura

        if (
            linha_fim_servidores
            >= linha_inicio_servidores
        ):
            chart = BarChart()
            chart.type = "bar"
            chart.style = 10
            chart.title = "Disponibilidade por MGW"
            chart.x_axis.title = "Disponibilidade"
            chart.y_axis.title = "Servidor"

            data = Reference(
                ws,
                min_col=2,
                min_row=linha_inicio_servidores - 1,
                max_row=linha_fim_servidores,
            )

            cats = Reference(
                ws,
                min_col=1,
                min_row=linha_inicio_servidores,
                max_row=linha_fim_servidores,
            )

            chart.add_data(
                data,
                titles_from_data=True,
            )
            chart.set_categories(cats)
            chart.height = 7
            chart.width = 12
            ws.add_chart(chart, "F8")

    def _criar_estatisticas(self, workbook, dados):
        ws = workbook.create_sheet("ESTATÍSTICAS")
        ws.sheet_view.showGridLines = False

        linhas = [
            (
                "Período diário",
                f"{dados['inicio_periodo']:%d/%m/%Y %H:%M:%S} a "
                f"{dados['fim_periodo']:%d/%m/%Y %H:%M:%S}",
            ),
            (
                "Período dos rankings",
                f"{dados['inicio_ranking']:%d/%m/%Y %H:%M:%S} a "
                f"{dados['fim_ranking']:%d/%m/%Y %H:%M:%S}",
            ),
            (
                "Quedas de rotas no dia",
                dados["total_quedas_rotas"],
            ),
            (
                "Recuperações de rotas no dia",
                dados["total_recuperacoes_rotas"],
            ),
        ]

        ws.append(["Indicador", "Valor"])

        for item in linhas:
            ws.append(list(item))

        self._formatar_tabela_simples(
            ws,
            "TabelaEstatisticas",
            2,
        )
        ws.column_dimensions["A"].width = 36
        ws.column_dimensions["B"].width = 55

    def _criar_enderecos(self, workbook, dados):
        ws = workbook.create_sheet("ENDEREÇOS")
        ws.sheet_view.showGridLines = False

        ws.append(
            [
                "Endereço",
                "Quedas (7 dias)",
                "Rotas associadas",
                "Servidores afetados",
                "Tempo total (7 dias)",
                "Média por queda",
                "Disponibilidade (7 dias)",
            ]
        )

        for item in sorted(
            dados["enderecos"],
            key=lambda x: (
                -x["quedas"],
                x["endereco"].casefold(),
            ),
        ):
            ws.append(
                [
                    item["endereco"],
                    item["quedas"],
                    "\n".join(item["rotas"]),
                    ", ".join(item["servidores"]),
                    self._formatar_duracao(
                        item["tempo_total_segundos"]
                    ),
                    self._formatar_duracao(
                        item["tempo_medio_segundos"]
                    ),
                    item["disponibilidade"] / 100,
                ]
            )
            ws.cell(
                ws.max_row,
                7,
            ).number_format = "0.00%"

        self._formatar_tabela_simples(
            ws,
            "TabelaEnderecos",
            7,
        )
        self._ajustar_larguras(
            ws,
            {
                1: 30,
                2: 18,
                3: 55,
                4: 35,
                5: 22,
                6: 20,
                7: 24,
            },
        )

    def _criar_rotas(self, workbook, dados):
        ws = workbook.create_sheet("ROTAS")
        ws.sheet_view.showGridLines = False

        ws.append(
            [
                "Rota",
                "Quedas (7 dias)",
                "Endereços",
                "Servidores",
                "Tempo total (7 dias)",
                "Média",
                "Disponibilidade (7 dias)",
            ]
        )

        for item in sorted(
            dados["rotas"],
            key=lambda x: (
                -x["quedas"],
                x["rota"].casefold(),
            ),
        ):
            ws.append(
                [
                    item["rota"],
                    item["quedas"],
                    "\n".join(item["enderecos"]),
                    ", ".join(item["servidores"]),
                    self._formatar_duracao(
                        item["tempo_total_segundos"]
                    ),
                    self._formatar_duracao(
                        item["tempo_medio_segundos"]
                    ),
                    item["disponibilidade"] / 100,
                ]
            )
            ws.cell(
                ws.max_row,
                7,
            ).number_format = "0.00%"

        self._formatar_tabela_simples(
            ws,
            "TabelaRotasDashboard",
            7,
        )
        self._ajustar_larguras(
            ws,
            {
                1: 55,
                2: 18,
                3: 35,
                4: 35,
                5: 22,
                6: 20,
                7: 24,
            },
        )

    def _criar_servidores(self, workbook, dados):
        ws = workbook.create_sheet("SERVIDORES")
        ws.sheet_view.showGridLines = False

        ws.append(
            [
                "Servidor",
                "Rotas monitoradas",
                "Quedas no dia",
                "Tempo indisponível no dia",
                "Indisponibilidade atual / média",
                "Disponibilidade diária",
            ]
        )

        for item in sorted(
            dados["servidores"],
            key=lambda x: (
                x["servidor"].casefold()
            ),
        ):
            if item["indisponivel_atual"]:
                referencia = self._formatar_duracao_horas(
                    item[
                        "indisponibilidade_atual_segundos"
                    ]
                )
            else:
                referencia = self._formatar_duracao(
                    item["tempo_medio_segundos"]
                )

            ws.append(
                [
                    item["servidor"],
                    item["rotas_monitoradas"],
                    item["quedas"],
                    self._formatar_duracao(
                        item["tempo_total_segundos"]
                    ),
                    referencia,
                    item["disponibilidade"] / 100,
                ]
            )
            ws.cell(
                ws.max_row,
                6,
            ).number_format = "0.00%"

        self._formatar_tabela_simples(
            ws,
            "TabelaServidoresDashboard",
            6,
        )
        self._ajustar_larguras(
            ws,
            {
                1: 18,
                2: 20,
                3: 18,
                4: 25,
                5: 32,
                6: 24,
            },
        )

    @staticmethod
    def _formatar_cabecalho(
        ws,
        linha,
        col_ini,
        col_fim,
    ):
        for col in range(col_ini, col_fim + 1):
            celula = ws.cell(linha, col)
            celula.fill = PatternFill(
                "solid",
                fgColor="37474F",
            )
            celula.font = Font(
                bold=True,
                color="FFFFFF",
            )
            celula.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True,
            )

    @staticmethod
    def _formatar_corpo(
        ws,
        linha_ini,
        linha_fim,
        col_ini,
        col_fim,
    ):
        if linha_ini > linha_fim:
            return

        lado = Side(
            style="thin",
            color="D9D9D9",
        )
        borda = Border(
            left=lado,
            right=lado,
            top=lado,
            bottom=lado,
        )

        for linha in ws.iter_rows(
            min_row=linha_ini,
            max_row=linha_fim,
            min_col=col_ini,
            max_col=col_fim,
        ):
            for celula in linha:
                celula.border = borda
                celula.alignment = Alignment(
                    vertical="center",
                    wrap_text=True,
                )

    def _formatar_tabela_simples(
        self,
        ws,
        nome_tabela,
        qtd_colunas,
    ):
        self._formatar_cabecalho(
            ws,
            1,
            1,
            qtd_colunas,
        )
        self._formatar_corpo(
            ws,
            2,
            ws.max_row,
            1,
            qtd_colunas,
        )
        ws.freeze_panes = "A2"

        if ws.max_row >= 2:
            tabela = Table(
                displayName=nome_tabela,
                ref=(
                    f"A1:"
                    f"{get_column_letter(qtd_colunas)}"
                    f"{ws.max_row}"
                ),
            )
            tabela.tableStyleInfo = TableStyleInfo(
                name="TableStyleMedium2",
                showFirstColumn=False,
                showLastColumn=False,
                showRowStripes=False,
                showColumnStripes=False,
            )
            ws.add_table(tabela)

    @staticmethod
    def _ajustar_larguras(ws, larguras):
        for numero_coluna, largura in larguras.items():
            ws.column_dimensions[
                get_column_letter(numero_coluna)
            ].width = largura

        for linha in ws.iter_rows(min_row=2):
            for celula in linha:
                celula.alignment = Alignment(
                    vertical="center",
                    wrap_text=True,
                )

    @classmethod
    def _complementar_eventos_atuais(
        cls,
        eventos,
        naps,
        banco,
        fim_periodo,
    ):
        """
        Complementa os eventos do Dashboard usando o histórico de
        coletas para rotas que continuam indisponíveis.

        Isso permite que um IP Peer que caiu em um dia anterior e
        continua indisponível seja considerado no ranking móvel de
        7 dias, mesmo quando o evento original não foi registrado na
        tabela de eventos no momento exato da queda.
        """
        if banco is None:
            return eventos

        resultado = list(eventos)

        for nap in naps:
            try:
                if not nap.esta_inativo():
                    continue
            except Exception:
                continue

            servidor = str(
                nap.servidor or "Não informado"
            ).strip()
            rota = str(
                nap.nome or "Rota não informada"
            ).strip()
            endereco = str(
                nap.endereco() or "Não informado"
            ).strip()

            try:
                inicio = banco.obter_inicio_indisponibilidade_atual(
                    servidor=servidor,
                    rota=rota,
                    endereco=(
                        ""
                        if endereco == "Não informado"
                        else endereco
                    ),
                )
            except Exception:
                inicio = None

            inicio = cls._converter_data(inicio)

            if inicio is None or inicio > fim_periodo:
                continue

            resultado.append(
                {
                    "servidor": servidor,
                    "rota": rota,
                    "endereco": endereco,
                    "caiu_em": inicio.isoformat(
                        timespec="seconds"
                    ),
                    "voltou_em": None,
                    "status": "ABERTO",
                    "origem": "historico_coletas",
                }
            )

        return resultado

    @classmethod
    def _quantidade_incidentes_uniao(
        cls,
        eventos,
        inicio_periodo,
        fim_periodo,
    ):
        """
        Conta incidentes do IP Peer unificando períodos sobrepostos
        de várias rotas e/ou MGWs do mesmo endereço.
        """
        intervalos = []

        for evento in eventos:
            caiu = cls._converter_data(
                evento.get("caiu_em")
            )
            voltou = cls._converter_data(
                evento.get("voltou_em")
            )

            if caiu is None or caiu > fim_periodo:
                continue

            fim = min(
                voltou or fim_periodo,
                fim_periodo,
            )

            if fim >= inicio_periodo:
                # Preserva o início real do incidente. Assim, uma
                # indisponibilidade que começou antes da janela de
                # 7 dias não é contada como uma nova queda apenas
                # por continuar aberta dentro da janela.
                intervalos.append((caiu, fim))

        if not intervalos:
            return 0

        intervalos.sort(key=lambda item: item[0])

        unidos = []
        inicio_atual, fim_atual = intervalos[0]

        for inicio, fim in intervalos[1:]:
            if inicio <= fim_atual:
                fim_atual = max(fim_atual, fim)
            else:
                unidos.append((inicio_atual, fim_atual))
                inicio_atual, fim_atual = inicio, fim

        unidos.append((inicio_atual, fim_atual))

        return sum(
            1
            for inicio, _ in unidos
            if inicio_periodo <= inicio <= fim_periodo
        )

    @classmethod
    def _eventos_sobrepostos(
        cls,
        eventos,
        inicio_periodo,
        fim_periodo,
    ):
        resultado = []

        for evento in eventos:
            caiu = cls._converter_data(
                evento.get("caiu_em")
            )
            voltou = cls._converter_data(
                evento.get("voltou_em")
            )

            if caiu is None or caiu > fim_periodo:
                continue

            if (
                voltou is not None
                and voltou < inicio_periodo
            ):
                continue

            resultado.append(evento)

        return resultado

    @classmethod
    def _duracao_sobreposta(
        cls,
        evento,
        inicio_periodo,
        fim_periodo,
    ):
        caiu = cls._converter_data(
            evento.get("caiu_em")
        )
        voltou = cls._converter_data(
            evento.get("voltou_em")
        )

        if caiu is None:
            return 0

        inicio = max(caiu, inicio_periodo)
        fim = min(
            voltou or fim_periodo,
            fim_periodo,
        )

        return max(
            0,
            int((fim - inicio).total_seconds()),
        )

    @classmethod
    def _duracao_uniao(
        cls,
        eventos,
        inicio_periodo,
        fim_periodo,
    ):
        intervalos = []

        for evento in eventos:
            caiu = cls._converter_data(
                evento.get("caiu_em")
            )
            voltou = cls._converter_data(
                evento.get("voltou_em")
            )

            if caiu is None:
                continue

            inicio = max(caiu, inicio_periodo)
            fim = min(
                voltou or fim_periodo,
                fim_periodo,
            )

            if fim > inicio:
                intervalos.append((inicio, fim))

        if not intervalos:
            return 0

        intervalos.sort(
            key=lambda item: item[0]
        )

        inicio_atual, fim_atual = intervalos[0]
        total = 0

        for inicio, fim in intervalos[1:]:
            if inicio <= fim_atual:
                fim_atual = max(
                    fim_atual,
                    fim,
                )
            else:
                total += int(
                    (
                        fim_atual - inicio_atual
                    ).total_seconds()
                )
                inicio_atual, fim_atual = (
                    inicio,
                    fim,
                )

        total += int(
            (
                fim_atual - inicio_atual
            ).total_seconds()
        )

        return total

    @classmethod
    def _esta_no_periodo(
        cls,
        valor,
        inicio_periodo,
        fim_periodo,
    ):
        data = cls._converter_data(valor)
        return (
            data is not None
            and inicio_periodo
            <= data
            <= fim_periodo
        )

    @classmethod
    def _evento_aberto_mais_antigo(
        cls,
        eventos,
    ):
        abertos = []

        for evento in eventos:
            voltou = cls._converter_data(
                evento.get("voltou_em")
            )
            status = str(
                evento.get("status") or ""
            ).upper()

            if voltou is None or status == "ABERTO":
                caiu = cls._converter_data(
                    evento.get("caiu_em")
                )

                if caiu is not None:
                    abertos.append((caiu, evento))

        if not abertos:
            return None

        abertos.sort(key=lambda item: item[0])
        return abertos[0][1]

    @classmethod
    def _duracao_evento_completa(cls, evento):
        caiu = cls._converter_data(
            evento.get("caiu_em")
        )
        voltou = cls._converter_data(
            evento.get("voltou_em")
        )

        if caiu is None or voltou is None:
            return None

        return max(
            0,
            int((voltou - caiu).total_seconds()),
        )

    @staticmethod
    def _converter_data(valor):
        if valor is None:
            return None

        if isinstance(valor, datetime):
            return valor

        try:
            return datetime.fromisoformat(
                str(valor)
            )
        except (ValueError, TypeError):
            return None

    @staticmethod
    def _formatar_duracao(segundos):
        segundos = max(
            0,
            int(segundos or 0),
        )

        dias, resto = divmod(
            segundos,
            86400,
        )
        horas, resto = divmod(
            resto,
            3600,
        )
        minutos, segundos = divmod(
            resto,
            60,
        )

        partes = []

        if dias:
            partes.append(
                f"{dias} dia"
                if dias == 1
                else f"{dias} dias"
            )

        if horas:
            partes.append(f"{horas} h")

        if minutos:
            partes.append(f"{minutos} min")

        if segundos or not partes:
            partes.append(f"{segundos} s")

        return " ".join(partes)

    @staticmethod
    def _formatar_duracao_horas(segundos):
        """
        Mostra indisponibilidades longas em horas acumuladas.
        Ex.: 14 dias = 336 h.
        """
        segundos = max(
            0,
            int(segundos or 0),
        )

        horas_totais, resto = divmod(
            segundos,
            3600,
        )
        minutos, segundos = divmod(
            resto,
            60,
        )

        partes = []

        if horas_totais:
            partes.append(
                f"{horas_totais} h"
            )

        if minutos:
            partes.append(
                f"{minutos} min"
            )

        if segundos and horas_totais < 24:
            partes.append(
                f"{segundos} s"
            )

        if not partes:
            partes.append("0 s")

        return " ".join(partes)

    @staticmethod
    def _formatar_percentual(valor):
        return f"{valor:.2f}%".replace(
            ".",
            ",",
        )

    @staticmethod
    def _medalha(posicao):
        return {
            1: "🥇",
            2: "🥈",
            3: "🥉",
        }.get(
            posicao,
            f"{posicao}º",
        )
