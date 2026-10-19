"""
MGW Monitor - Histórico consolidado de congestionamentos do Portal Orbe.

Fonte de verdade:
    dados/portal_orbe/portal_orbe_YYYY-MM-DD.jsonl

Saídas diárias:
    dados/portal_orbe/eventos_consolidados/
        portal_orbe_eventos_YYYY-MM-DD.jsonl

    relatorios/portal_orbe/
        portal_orbe_congestionamentos_YYYY-MM-DD.xlsx

Princípios:
- NÃO altera o JSONL bruto.
- NÃO envia alertas Telegram em tempo real.
- Eventos transitórios são derivados do histórico bruto.
- Incidentes confirmados são derivados das transições
  incidente_confirmado=False -> True.
- Recorrência confirmada é a 2ª, 3ª... confirmação da mesma fila no mesmo dia.
- A geração é idempotente: reprocessar o dia sobrescreve o consolidado do dia,
  evitando duplicidades.
"""

from __future__ import annotations

import json
from collections import defaultdict
from dataclasses import asdict, dataclass
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Tuple

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError:
    Workbook = None


PASTA_BRUTA_PADRAO = Path("dados") / "portal_orbe"
PASTA_CONSOLIDADA_PADRAO = (
    PASTA_BRUTA_PADRAO / "eventos_consolidados"
)
PASTA_RELATORIOS_PADRAO = Path("relatorios") / "portal_orbe"

LIMIAR_ALTA_RECORRENCIA_TRANSITORIA_PADRAO = 5
MAX_GAP_COLETA_MINUTOS_PADRAO = 10


def _dt(valor: Any) -> Optional[datetime]:
    if not valor:
        return None
    if isinstance(valor, datetime):
        return valor
    try:
        return datetime.fromisoformat(str(valor))
    except (TypeError, ValueError):
        return None


def _float(valor: Any, padrao: float = 0.0) -> float:
    try:
        return float(valor)
    except (TypeError, ValueError):
        return float(padrao)


def _int(valor: Any, padrao: int = 0) -> int:
    try:
        return int(valor)
    except (TypeError, ValueError):
        return int(padrao)


def _fmt_duracao(segundos: float) -> str:
    segundos = max(0, int(round(segundos)))
    horas, resto = divmod(segundos, 3600)
    minutos, seg = divmod(resto, 60)

    if horas:
        return f"{horas}h{minutos:02d}min"
    if minutos:
        return f"{minutos}min{seg:02d}s"
    return f"{seg}s"


def _chave(amostra: Dict[str, Any]) -> str:
    return (
        f"inst{amostra.get('instancia')}|"
        f"{amostra.get('projeto_id')}|"
        f"{amostra.get('codigo')}"
    )


def _acima_p34(amostra: Dict[str, Any]) -> bool:
    status = amostra.get("status34") or {}
    return (
        _float(status.get("percentual"))
        > _float(amostra.get("limite_p34"), 20.0)
    )


def _acima_p135(amostra: Dict[str, Any]) -> bool:
    status = amostra.get("status135") or {}
    return (
        _float(status.get("percentual"))
        > _float(amostra.get("limite_p135"), 20.0)
    )


def _classificacao(p34: bool, p135: bool) -> str:
    if p34 and p135:
        return (
            "Congestionamento na Companhia (P34) + "
            "Congestionamento Interno (P135)"
        )
    if p34:
        return "Congestionamento na Companhia (P34)"
    if p135:
        return "Congestionamento Interno (P135)"
    return "Sem ultrapassagem"


@dataclass
class EpisodioAcima:
    chave: str
    instancia: str
    projeto_id: str
    projeto: str
    fila: str
    codigo: str
    inicio: datetime
    fim: datetime
    ultima_acima: datetime
    p34_pico: float
    p34_qtd_pico: int
    p135_pico: float
    p135_qtd_pico: int
    teve_p34: bool
    teve_p135: bool
    limite_p34: float
    limite_p135: float
    qualidade: str = "OK"

    def contem(self, momento: datetime) -> bool:
        return self.inicio <= momento <= self.fim


@dataclass
class EventoConsolidado:
    data: str
    tipo_evento: str
    instancia: str
    projeto_id: str
    projeto: str
    fila: str
    codigo: str

    inicio: str
    fim: Optional[str]
    duracao_segundos: float
    duracao: str

    classificacao: str
    p34_pico: float
    p34_qtd_pico: int
    p135_pico: float
    p135_qtd_pico: int
    limite_p34: float
    limite_p135: float

    confirmado_em: Optional[str] = None
    normalizado_em: Optional[str] = None
    ocorrencia_confirmada_dia: Optional[int] = None

    qualidade: str = "OK"
    observacao: str = ""

    def dict(self) -> Dict[str, Any]:
        return asdict(self)


def caminho_bruto(
    data_ref: date,
    pasta_bruta: Path = PASTA_BRUTA_PADRAO,
) -> Path:
    return (
        Path(pasta_bruta)
        / f"portal_orbe_{data_ref:%Y-%m-%d}.jsonl"
    )


def caminho_consolidado(
    data_ref: date,
    pasta_consolidada: Path = PASTA_CONSOLIDADA_PADRAO,
) -> Path:
    return (
        Path(pasta_consolidada)
        / f"portal_orbe_eventos_{data_ref:%Y-%m-%d}.jsonl"
    )


def caminho_excel(
    data_ref: date,
    pasta_relatorios: Path = PASTA_RELATORIOS_PADRAO,
) -> Path:
    return (
        Path(pasta_relatorios)
        / f"portal_orbe_congestionamentos_{data_ref:%Y-%m-%d}.xlsx"
    )


def carregar_amostras(caminho: Path) -> List[Dict[str, Any]]:
    if not caminho.exists():
        raise FileNotFoundError(
            f"Histórico bruto não encontrado: {caminho}"
        )

    saida: List[Dict[str, Any]] = []

    with caminho.open("r", encoding="utf-8") as arquivo:
        for numero_linha, linha in enumerate(arquivo, start=1):
            linha = linha.strip()
            if not linha:
                continue

            try:
                item = json.loads(linha)
            except json.JSONDecodeError:
                continue

            momento = _dt(item.get("timestamp"))
            if momento is None:
                continue

            item["_dt"] = momento
            item["_linha"] = numero_linha
            saida.append(item)

    saida.sort(
        key=lambda x: (
            x["_dt"],
            str(x.get("instancia", "")),
            str(x.get("codigo", "")),
        )
    )
    return saida


def _novo_episodio(
    chave: str,
    item: Dict[str, Any],
) -> Dict[str, Any]:
    status34 = item.get("status34") or {}
    status135 = item.get("status135") or {}

    return {
        "chave": chave,
        "instancia": str(item.get("instancia", "")),
        "projeto_id": str(item.get("projeto_id", "")),
        "projeto": str(item.get("projeto", "")),
        "fila": str(item.get("fila", "")),
        "codigo": str(item.get("codigo", "")),
        "inicio": item["_dt"],
        "ultima_acima": item["_dt"],
        "p34_pico": _float(status34.get("percentual")),
        "p34_qtd_pico": _int(status34.get("quantidade")),
        "p135_pico": _float(status135.get("percentual")),
        "p135_qtd_pico": _int(status135.get("quantidade")),
        "teve_p34": _acima_p34(item),
        "teve_p135": _acima_p135(item),
        "limite_p34": _float(item.get("limite_p34"), 20.0),
        "limite_p135": _float(item.get("limite_p135"), 20.0),
    }


def _atualizar_episodio(
    controle: Dict[str, Any],
    item: Dict[str, Any],
) -> None:
    status34 = item.get("status34") or {}
    status135 = item.get("status135") or {}

    p34 = _float(status34.get("percentual"))
    p135 = _float(status135.get("percentual"))

    if p34 > controle["p34_pico"]:
        controle["p34_pico"] = p34
        controle["p34_qtd_pico"] = _int(
            status34.get("quantidade")
        )

    if p135 > controle["p135_pico"]:
        controle["p135_pico"] = p135
        controle["p135_qtd_pico"] = _int(
            status135.get("quantidade")
        )

    controle["teve_p34"] = (
        controle["teve_p34"] or _acima_p34(item)
    )
    controle["teve_p135"] = (
        controle["teve_p135"] or _acima_p135(item)
    )
    controle["ultima_acima"] = item["_dt"]


def extrair_episodios_acima(
    amostras: Iterable[Dict[str, Any]],
    *,
    max_gap_coleta_minutos: int = (
        MAX_GAP_COLETA_MINUTOS_PADRAO
    ),
) -> List[EpisodioAcima]:
    """
    Constrói todas as janelas contínuas acima de P34/P135
    dentro do horário operacional.

    Se houver lacuna de coleta maior que max_gap_coleta_minutos,
    o episódio é encerrado e marcado como LACUNA_COLETA.
    """
    por_fila: Dict[str, List[Dict[str, Any]]] = defaultdict(list)
    for item in amostras:
        por_fila[_chave(item)].append(item)

    max_gap = timedelta(minutes=max_gap_coleta_minutos)
    episodios: List[EpisodioAcima] = []

    for chave, itens in por_fila.items():
        itens.sort(key=lambda x: x["_dt"])

        atual: Optional[Dict[str, Any]] = None
        anterior: Optional[Dict[str, Any]] = None

        def fechar(
            fim: datetime,
            qualidade: str = "OK",
        ) -> None:
            nonlocal atual
            if atual is None:
                return

            episodios.append(
                EpisodioAcima(
                    chave=atual["chave"],
                    instancia=atual["instancia"],
                    projeto_id=atual["projeto_id"],
                    projeto=atual["projeto"],
                    fila=atual["fila"],
                    codigo=atual["codigo"],
                    inicio=atual["inicio"],
                    fim=fim,
                    ultima_acima=atual["ultima_acima"],
                    p34_pico=atual["p34_pico"],
                    p34_qtd_pico=atual["p34_qtd_pico"],
                    p135_pico=atual["p135_pico"],
                    p135_qtd_pico=atual["p135_qtd_pico"],
                    teve_p34=atual["teve_p34"],
                    teve_p135=atual["teve_p135"],
                    limite_p34=atual["limite_p34"],
                    limite_p135=atual["limite_p135"],
                    qualidade=qualidade,
                )
            )
            atual = None

        for item in itens:
            momento = item["_dt"]

            if (
                anterior is not None
                and momento - anterior["_dt"] > max_gap
            ):
                if atual is not None:
                    fechar(
                        atual["ultima_acima"],
                        "LACUNA_COLETA",
                    )

            operacional = bool(
                item.get(
                    "dentro_horario_operacional",
                    True,
                )
            )
            acima = bool(item.get("acima_limite")) and operacional

            if acima:
                if atual is None:
                    atual = _novo_episodio(chave, item)
                else:
                    _atualizar_episodio(atual, item)
            elif atual is not None:
                # O primeiro registro abaixo do limite define o fim
                # da janela observada.
                fechar(momento)

            anterior = item

        if atual is not None:
            fechar(
                atual["ultima_acima"],
                "ARQUIVO_EM_ABERTO",
            )

    episodios.sort(
        key=lambda x: (
            x.inicio,
            x.instancia,
            x.codigo,
        )
    )
    return episodios


def extrair_transicoes_confirmacao(
    amostras: Iterable[Dict[str, Any]],
) -> Dict[str, List[Dict[str, Any]]]:
    """
    Retorna as transições reais do estado persistido pelo monitor:
    False -> True = confirmação
    True -> False = normalização
    """
    por_fila: Dict[str, List[Dict[str, Any]]] = defaultdict(list)
    for item in amostras:
        por_fila[_chave(item)].append(item)

    resultado: Dict[str, List[Dict[str, Any]]] = {}

    for chave, itens in por_fila.items():
        itens.sort(key=lambda x: x["_dt"])
        anterior = False
        eventos: List[Dict[str, Any]] = []

        for item in itens:
            atual = bool(item.get("incidente_confirmado"))

            if atual and not anterior:
                eventos.append({
                    "tipo": "CONFIRMACAO",
                    "momento": item["_dt"],
                    "amostra": item,
                })
            elif anterior and not atual:
                eventos.append({
                    "tipo": "NORMALIZACAO",
                    "momento": item["_dt"],
                    "amostra": item,
                })

            anterior = atual

        resultado[chave] = eventos

    return resultado


def _episodio_da_confirmacao(
    episodios: List[EpisodioAcima],
    momento: datetime,
) -> Optional[EpisodioAcima]:
    candidatos = [
        episodio
        for episodio in episodios
        if episodio.inicio <= momento <= episodio.fim
    ]

    if candidatos:
        return max(candidatos, key=lambda x: x.inicio)

    # Segurança para pequena diferença entre a última amostra
    # acima do limite e a amostra que gravou a confirmação.
    anteriores = [
        episodio
        for episodio in episodios
        if episodio.inicio <= momento
        and momento - episodio.fim <= timedelta(minutes=2)
    ]
    if anteriores:
        return max(anteriores, key=lambda x: x.inicio)

    return None


def reconstruir_eventos(
    amostras: List[Dict[str, Any]],
    *,
    max_gap_coleta_minutos: int = (
        MAX_GAP_COLETA_MINUTOS_PADRAO
    ),
) -> Tuple[
    List[EventoConsolidado],
    List[EventoConsolidado],
]:
    episodios = extrair_episodios_acima(
        amostras,
        max_gap_coleta_minutos=max_gap_coleta_minutos,
    )
    transicoes = extrair_transicoes_confirmacao(amostras)

    episodios_por_chave: Dict[str, List[EpisodioAcima]] = (
        defaultdict(list)
    )
    for episodio in episodios:
        episodios_por_chave[episodio.chave].append(episodio)

    episodios_confirmados_ids = set()
    confirmados: List[EventoConsolidado] = []

    for chave, eventos_estado in transicoes.items():
        numero = 0
        aberto: Optional[EventoConsolidado] = None

        for evento_estado in eventos_estado:
            tipo = evento_estado["tipo"]
            momento = evento_estado["momento"]
            amostra = evento_estado["amostra"]

            if tipo == "CONFIRMACAO":
                numero += 1

                episodio = _episodio_da_confirmacao(
                    episodios_por_chave.get(chave, []),
                    momento,
                )

                if episodio is not None:
                    episodios_confirmados_ids.add(id(episodio))
                    inicio = episodio.inicio
                    p34_pico = episodio.p34_pico
                    p34_qtd = episodio.p34_qtd_pico
                    p135_pico = episodio.p135_pico
                    p135_qtd = episodio.p135_qtd_pico
                    limite34 = episodio.limite_p34
                    limite135 = episodio.limite_p135
                    classif = _classificacao(
                        episodio.teve_p34,
                        episodio.teve_p135,
                    )
                    qualidade = episodio.qualidade
                    projeto = episodio.projeto
                    projeto_id = episodio.projeto_id
                    fila = episodio.fila
                    codigo = episodio.codigo
                    instancia = episodio.instancia
                else:
                    inicio = momento
                    status34 = amostra.get("status34") or {}
                    status135 = amostra.get("status135") or {}
                    p34_pico = _float(status34.get("percentual"))
                    p34_qtd = _int(status34.get("quantidade"))
                    p135_pico = _float(status135.get("percentual"))
                    p135_qtd = _int(status135.get("quantidade"))
                    limite34 = _float(
                        amostra.get("limite_p34"),
                        20.0,
                    )
                    limite135 = _float(
                        amostra.get("limite_p135"),
                        20.0,
                    )
                    classif = _classificacao(
                        _acima_p34(amostra),
                        _acima_p135(amostra),
                    )
                    qualidade = "SEM_EPISODIO_ACIMA_LOCALIZADO"
                    projeto = str(amostra.get("projeto", ""))
                    projeto_id = str(
                        amostra.get("projeto_id", "")
                    )
                    fila = str(amostra.get("fila", ""))
                    codigo = str(amostra.get("codigo", ""))
                    instancia = str(
                        amostra.get("instancia", "")
                    )

                aberto = EventoConsolidado(
                    data=momento.date().isoformat(),
                    tipo_evento=(
                        "CONFIRMADO"
                        if numero == 1
                        else "RECORRENCIA_CONFIRMADA"
                    ),
                    instancia=instancia,
                    projeto_id=projeto_id,
                    projeto=projeto,
                    fila=fila,
                    codigo=codigo,
                    inicio=inicio.isoformat(
                        timespec="seconds"
                    ),
                    fim=None,
                    duracao_segundos=0.0,
                    duracao="Em aberto",
                    classificacao=classif,
                    p34_pico=p34_pico,
                    p34_qtd_pico=p34_qtd,
                    p135_pico=p135_pico,
                    p135_qtd_pico=p135_qtd,
                    limite_p34=limite34,
                    limite_p135=limite135,
                    confirmado_em=momento.isoformat(
                        timespec="seconds"
                    ),
                    ocorrencia_confirmada_dia=numero,
                    qualidade=qualidade,
                    observacao=(
                        "Primeiro incidente confirmado da fila no dia."
                        if numero == 1
                        else (
                            f"Recorrência confirmada nº {numero} "
                            "da mesma fila no dia."
                        )
                    ),
                )
                confirmados.append(aberto)

            elif tipo == "NORMALIZACAO" and aberto is not None:
                aberto.normalizado_em = momento.isoformat(
                    timespec="seconds"
                )
                aberto.fim = aberto.normalizado_em
                inicio_dt = _dt(aberto.inicio) or momento
                duracao = max(
                    0.0,
                    (momento - inicio_dt).total_seconds(),
                )
                aberto.duracao_segundos = duracao
                aberto.duracao = _fmt_duracao(duracao)
                aberto = None

    transitorios: List[EventoConsolidado] = []

    for episodio in episodios:
        if id(episodio) in episodios_confirmados_ids:
            continue

        duracao = max(
            0.0,
            (episodio.fim - episodio.inicio).total_seconds(),
        )

        observacao = (
            "Ultrapassagem de limite sem confirmação do incidente."
        )
        if episodio.qualidade == "ARQUIVO_EM_ABERTO":
            observacao = (
                "Episódio ainda estava acima do limite no final do arquivo. "
                "Reprocessar após o fechamento do dia."
            )
        elif episodio.qualidade == "LACUNA_COLETA":
            observacao = (
                "Episódio interrompido por lacuna de coleta; preservado "
                "historicamente para auditoria."
            )

        transitorios.append(
            EventoConsolidado(
                data=episodio.inicio.date().isoformat(),
                tipo_evento="TRANSITORIO",
                instancia=episodio.instancia,
                projeto_id=episodio.projeto_id,
                projeto=episodio.projeto,
                fila=episodio.fila,
                codigo=episodio.codigo,
                inicio=episodio.inicio.isoformat(
                    timespec="seconds"
                ),
                fim=episodio.fim.isoformat(
                    timespec="seconds"
                ),
                duracao_segundos=duracao,
                duracao=_fmt_duracao(duracao),
                classificacao=_classificacao(
                    episodio.teve_p34,
                    episodio.teve_p135,
                ),
                p34_pico=episodio.p34_pico,
                p34_qtd_pico=episodio.p34_qtd_pico,
                p135_pico=episodio.p135_pico,
                p135_qtd_pico=episodio.p135_qtd_pico,
                limite_p34=episodio.limite_p34,
                limite_p135=episodio.limite_p135,
                qualidade=episodio.qualidade,
                observacao=observacao,
            )
        )

    confirmados.sort(
        key=lambda x: (
            x.inicio,
            x.instancia,
            x.codigo,
        )
    )
    transitorios.sort(
        key=lambda x: (
            x.inicio,
            x.instancia,
            x.codigo,
        )
    )

    return confirmados, transitorios


def calcular_recorrencia_transitoria(
    transitorios: List[EventoConsolidado],
    confirmados: List[EventoConsolidado],
    *,
    limiar_alta_recorrencia: int = (
        LIMIAR_ALTA_RECORRENCIA_TRANSITORIA_PADRAO
    ),
) -> List[Dict[str, Any]]:
    por_fila: Dict[
        Tuple[str, str, str, str],
        List[EventoConsolidado],
    ] = defaultdict(list)

    confirmados_por_fila: Dict[
        Tuple[str, str, str, str],
        int,
    ] = defaultdict(int)

    for evento in transitorios:
        chave = (
            evento.instancia,
            evento.projeto_id,
            evento.fila,
            evento.codigo,
        )
        por_fila[chave].append(evento)

    for evento in confirmados:
        chave = (
            evento.instancia,
            evento.projeto_id,
            evento.fila,
            evento.codigo,
        )
        confirmados_por_fila[chave] += 1

    linhas: List[Dict[str, Any]] = []

    for chave, eventos in por_fila.items():
        eventos.sort(key=lambda x: x.inicio)

        instancia, projeto_id, fila, codigo = chave
        maior_duracao = max(
            (
                evento.duracao_segundos
                for evento in eventos
            ),
            default=0.0,
        )

        linhas.append({
            "instancia": instancia,
            "projeto_id": projeto_id,
            "projeto": eventos[0].projeto,
            "fila": fila,
            "codigo": codigo,
            "eventos_transitorios": len(eventos),
            "incidentes_confirmados": (
                confirmados_por_fila.get(chave, 0)
            ),
            "p34_pico_transitorio": max(
                (x.p34_pico for x in eventos),
                default=0.0,
            ),
            "p135_pico_transitorio": max(
                (x.p135_pico for x in eventos),
                default=0.0,
            ),
            "maior_duracao_transitoria_seg": maior_duracao,
            "maior_duracao_transitoria": _fmt_duracao(
                maior_duracao
            ),
            "primeiro_transitorio": eventos[0].inicio,
            "ultimo_transitorio": eventos[-1].inicio,
            "classificacoes": " | ".join(
                sorted(
                    {
                        x.classificacao
                        for x in eventos
                    }
                )
            ),
            "alta_recorrencia": (
                len(eventos)
                >= limiar_alta_recorrencia
            ),
            "limiar_alta_recorrencia": (
                limiar_alta_recorrencia
            ),
        })

    linhas.sort(
        key=lambda x: (
            not x["alta_recorrencia"],
            -x["eventos_transitorios"],
            x["instancia"],
            x["fila"],
        )
    )
    return linhas


def salvar_consolidado_jsonl(
    eventos: List[EventoConsolidado],
    caminho: Path,
) -> Path:
    caminho.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    temporario = caminho.with_suffix(
        caminho.suffix + ".tmp"
    )

    with temporario.open(
        "w",
        encoding="utf-8",
    ) as arquivo:
        for evento in eventos:
            arquivo.write(
                json.dumps(
                    evento.dict(),
                    ensure_ascii=False,
                )
                + "\n"
            )

    temporario.replace(caminho)
    return caminho


def _cabecalhos_eventos() -> List[str]:
    return [
        "Tipo do evento",
        "Instância",
        "Projeto",
        "Fila",
        "Código",
        "Início",
        "Fim/Normalização",
        "Duração",
        "Classificação",
        "P34 pico (%)",
        "P34 chamadas",
        "P135 pico (%)",
        "P135 chamadas",
        "Limite P34 (%)",
        "Limite P135 (%)",
        "Confirmado em",
        "Ocorrência confirmada no dia",
        "Qualidade",
        "Observação",
    ]


def _linha_evento(
    evento: EventoConsolidado,
) -> List[Any]:
    return [
        evento.tipo_evento,
        evento.instancia,
        evento.projeto,
        evento.fila,
        evento.codigo,
        evento.inicio.replace("T", " "),
        (
            evento.normalizado_em
            or evento.fim
            or ""
        ).replace("T", " "),
        evento.duracao,
        evento.classificacao,
        evento.p34_pico,
        evento.p34_qtd_pico,
        evento.p135_pico,
        evento.p135_qtd_pico,
        evento.limite_p34,
        evento.limite_p135,
        (
            evento.confirmado_em.replace("T", " ")
            if evento.confirmado_em
            else ""
        ),
        evento.ocorrencia_confirmada_dia or "",
        evento.qualidade,
        evento.observacao,
    ]


def gerar_excel(
    *,
    data_ref: date,
    confirmados: List[EventoConsolidado],
    transitorios: List[EventoConsolidado],
    recorrencia: List[Dict[str, Any]],
    caminho: Path,
) -> Path:
    if Workbook is None:
        raise RuntimeError(
            "openpyxl não está instalado. "
            "Execute: pip install openpyxl"
        )

    caminho.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    wb = Workbook()
    ws_resumo = wb.active
    ws_resumo.title = "Resumo"

    todos = sorted(
        [*confirmados, *transitorios],
        key=lambda x: (
            x.inicio,
            x.instancia,
            x.codigo,
        ),
    )

    altas = [
        item
        for item in recorrencia
        if item["alta_recorrencia"]
    ]

    recorrencias_confirmadas = sum(
        1
        for item in confirmados
        if item.tipo_evento
        == "RECORRENCIA_CONFIRMADA"
    )

    linhas_resumo = [
        [
            "PORTAL ORBE — CONSOLIDADO DIÁRIO "
            "DE CONGESTIONAMENTOS"
        ],
        ["Data", data_ref.strftime("%d/%m/%Y")],
        [],
        ["Indicador", "Quantidade"],
        [
            "Incidentes confirmados",
            len(confirmados),
        ],
        [
            "Recorrências confirmadas",
            recorrencias_confirmadas,
        ],
        [
            "Eventos transitórios",
            len(transitorios),
        ],
        [
            "Filas com alta recorrência transitória",
            len(altas),
        ],
        [
            "Filas distintas com incidente confirmado",
            len(
                {
                    (x.instancia, x.codigo)
                    for x in confirmados
                }
            ),
        ],
        [
            "Filas distintas com transitórios",
            len(
                {
                    (x.instancia, x.codigo)
                    for x in transitorios
                }
            ),
        ],
        [],
        [
            "Critério de alta recorrência transitória",
            (
                f">= {recorrencia[0]['limiar_alta_recorrencia']} "
                "eventos/dia"
                if recorrencia
                else (
                    f">= "
                    f"{LIMIAR_ALTA_RECORRENCIA_TRANSITORIA_PADRAO} "
                    "eventos/dia"
                )
            ),
        ],
        [
            "Fonte",
            (
                "dados/portal_orbe/"
                f"portal_orbe_{data_ref:%Y-%m-%d}.jsonl"
            ),
        ],
    ]

    for linha in linhas_resumo:
        ws_resumo.append(linha)

    ws_eventos = wb.create_sheet("Eventos do Dia")
    ws_eventos.append(_cabecalhos_eventos())
    for evento in todos:
        ws_eventos.append(_linha_evento(evento))

    ws_conf = wb.create_sheet(
        "Incidentes Confirmados"
    )
    ws_conf.append(_cabecalhos_eventos())
    for evento in confirmados:
        ws_conf.append(_linha_evento(evento))

    ws_trans = wb.create_sheet(
        "Eventos Transitórios"
    )
    ws_trans.append(_cabecalhos_eventos())
    for evento in transitorios:
        ws_trans.append(_linha_evento(evento))

    ws_rec = wb.create_sheet(
        "Recorrência Transitória"
    )
    ws_rec.append([
        "Instância",
        "Projeto",
        "Fila",
        "Código",
        "Eventos transitórios",
        "Incidentes confirmados",
        "P34 pico transitório (%)",
        "P135 pico transitório (%)",
        "Maior duração transitória",
        "Primeiro transitório",
        "Último transitório",
        "Classificações",
        "Alta recorrência?",
        "Limiar",
    ])

    for item in recorrencia:
        ws_rec.append([
            item["instancia"],
            item["projeto"],
            item["fila"],
            item["codigo"],
            item["eventos_transitorios"],
            item["incidentes_confirmados"],
            item["p34_pico_transitorio"],
            item["p135_pico_transitorio"],
            item["maior_duracao_transitoria"],
            (
                item["primeiro_transitorio"]
                or ""
            ).replace("T", " "),
            (
                item["ultimo_transitorio"]
                or ""
            ).replace("T", " "),
            item["classificacoes"],
            (
                "SIM"
                if item["alta_recorrencia"]
                else "NÃO"
            ),
            item["limiar_alta_recorrencia"],
        ])

    cor_header = "1F4E78"
    cor_titulo = "0F243E"

    for ws in wb.worksheets:
        ws.freeze_panes = "A2"

        if ws.max_row >= 1:
            for celula in ws[1]:
                celula.font = Font(
                    bold=True,
                    color="FFFFFF",
                )
                celula.fill = PatternFill(
                    "solid",
                    fgColor=(
                        cor_titulo
                        if ws.title == "Resumo"
                        else cor_header
                    ),
                )
                celula.alignment = Alignment(
                    horizontal="center",
                    vertical="center",
                    wrap_text=True,
                )

        for coluna in ws.columns:
            letra = get_column_letter(
                coluna[0].column
            )
            maior = 0

            for celula in coluna[
                : min(ws.max_row, 400)
            ]:
                valor = (
                    ""
                    if celula.value is None
                    else str(celula.value)
                )
                maior = max(
                    maior,
                    len(valor),
                )

            ws.column_dimensions[letra].width = min(
                max(maior + 2, 11),
                42,
            )

        for linha in ws.iter_rows():
            for celula in linha:
                celula.alignment = Alignment(
                    vertical="top",
                    wrap_text=True,
                )

    for ws in (
        ws_eventos,
        ws_conf,
        ws_trans,
    ):
        for coluna in ("J", "L", "N", "O"):
            for celula in ws[coluna][1:]:
                celula.number_format = "0.00"

    for coluna in ("G", "H"):
        for celula in ws_rec[coluna][1:]:
            celula.number_format = "0.00"

    wb.save(caminho)
    return caminho


def montar_resumo_telegram(
    *,
    data_ref: date,
    confirmados: List[EventoConsolidado],
    transitorios: List[EventoConsolidado],
    recorrencia: List[Dict[str, Any]],
) -> str:
    altas = [
        item
        for item in recorrencia
        if item["alta_recorrencia"]
    ]

    recorrencias_confirmadas = sum(
        1
        for evento in confirmados
        if evento.tipo_evento
        == "RECORRENCIA_CONFIRMADA"
    )

    linhas = [
        "📊 <b>PORTAL ORBE — RESUMO DIÁRIO "
        "DE CONGESTIONAMENTOS</b>",
        "",
        f"📅 <b>Data:</b> {data_ref:%d/%m/%Y}",
        "",
        "📌 <b>Visão geral</b>",
        (
            "• Incidentes confirmados: "
            f"<b>{len(confirmados)}</b>"
        ),
        (
            "• Recorrências confirmadas da mesma fila: "
            f"<b>{recorrencias_confirmadas}</b>"
        ),
        (
            "• Eventos transitórios: "
            f"<b>{len(transitorios)}</b>"
        ),
        (
            "• Filas com alta recorrência transitória: "
            f"<b>{len(altas)}</b>"
        ),
        "",
        (
            "📎 O Excel anexo contém Instância, Projeto, Fila, "
            "horários, duração, picos de Congestionamento na "
            "Companhia (P34), Congestionamento Interno (P135), "
            "incidentes confirmados, eventos transitórios e "
            "recorrência transitória."
        ),
        "",
        (
            "ℹ️ Eventos transitórios não geram alerta individual "
            "no Telegram nem chamado Ticketdesk."
        ),
    ]

    if altas:
        linhas.extend([
            "",
            "🔎 <b>Filas com maior recorrência transitória</b>",
        ])

        for item in altas[:5]:
            linhas.append(
                f"• Instância {item['instancia']} — "
                f"{item['fila']} ({item['codigo']}): "
                f"<b>{item['eventos_transitorios']}</b> eventos"
            )

    return "\n".join(linhas)


def consolidar_dia(
    data_ref: date,
    *,
    pasta_bruta: Path = PASTA_BRUTA_PADRAO,
    pasta_consolidada: Path = (
        PASTA_CONSOLIDADA_PADRAO
    ),
    pasta_relatorios: Path = (
        PASTA_RELATORIOS_PADRAO
    ),
    limiar_alta_recorrencia: int = (
        LIMIAR_ALTA_RECORRENCIA_TRANSITORIA_PADRAO
    ),
    max_gap_coleta_minutos: int = (
        MAX_GAP_COLETA_MINUTOS_PADRAO
    ),
    gerar_xlsx: bool = True,
) -> Dict[str, Any]:
    bruto = caminho_bruto(
        data_ref,
        pasta_bruta,
    )

    amostras = carregar_amostras(bruto)

    confirmados, transitorios = reconstruir_eventos(
        amostras,
        max_gap_coleta_minutos=(
            max_gap_coleta_minutos
        ),
    )

    recorrencia = calcular_recorrencia_transitoria(
        transitorios,
        confirmados,
        limiar_alta_recorrencia=(
            limiar_alta_recorrencia
        ),
    )

    todos = sorted(
        [*confirmados, *transitorios],
        key=lambda x: (
            x.inicio,
            x.instancia,
            x.codigo,
        ),
    )

    arquivo_consolidado = caminho_consolidado(
        data_ref,
        pasta_consolidada,
    )
    salvar_consolidado_jsonl(
        todos,
        arquivo_consolidado,
    )

    arquivo_excel = None

    if gerar_xlsx:
        arquivo_excel = caminho_excel(
            data_ref,
            pasta_relatorios,
        )

        gerar_excel(
            data_ref=data_ref,
            confirmados=confirmados,
            transitorios=transitorios,
            recorrencia=recorrencia,
            caminho=arquivo_excel,
        )

    resumo = montar_resumo_telegram(
        data_ref=data_ref,
        confirmados=confirmados,
        transitorios=transitorios,
        recorrencia=recorrencia,
    )

    return {
        "data": data_ref.isoformat(),
        "arquivo_bruto": str(bruto),
        "arquivo_consolidado": str(
            arquivo_consolidado
        ),
        "arquivo_excel": (
            str(arquivo_excel)
            if arquivo_excel
            else None
        ),
        "confirmados": len(confirmados),
        "recorrencias_confirmadas": sum(
            1
            for evento in confirmados
            if evento.tipo_evento
            == "RECORRENCIA_CONFIRMADA"
        ),
        "transitorios": len(transitorios),
        "filas_alta_recorrencia_transitoria": (
            sum(
                1
                for item in recorrencia
                if item["alta_recorrencia"]
            )
        ),
        "resumo_telegram": resumo,
        "eventos_confirmados": [
            evento.dict()
            for evento in confirmados
        ],
        "eventos_transitorios": [
            evento.dict()
            for evento in transitorios
        ],
        "recorrencia_transitoria": recorrencia,
    }
