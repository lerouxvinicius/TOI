from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


class RelatorioExcel:
    def __init__(self, pasta_saida="relatorios"):
        self.pasta_saida = Path(pasta_saida)

    def gerar(self, naps, nomes_servidores):
        self.pasta_saida.mkdir(
            parents=True,
            exist_ok=True,
        )

        momento_coleta = datetime.now()

        data_hora_arquivo = momento_coleta.strftime(
            "%Y-%m-%d_%H-%M-%S"
        )

        caminho = (
            self.pasta_saida
            / f"relatorio_naps_{data_hora_arquivo}.xlsx"
        )

        workbook = Workbook()

        # Remove a aba vazia criada automaticamente.
        workbook.remove(workbook.active)

        naps_por_servidor = {
            nome_servidor: []
            for nome_servidor in nomes_servidores
        }

        for nap in naps:
            naps_por_servidor.setdefault(
                nap.servidor,
                [],
            ).append(nap)

        self._criar_aba_resumo(
            workbook=workbook,
            naps_por_servidor=naps_por_servidor,
            nomes_servidores=nomes_servidores,
            data_hora=momento_coleta,
        )

        self._criar_aba_links_consolidado(
            workbook=workbook,
            naps=naps,
        )

        self._criar_aba_inativos(
            workbook=workbook,
            naps=naps,
        )

        for indice, nome_servidor in enumerate(
            nomes_servidores,
            start=1,
        ):
            planilha = workbook.create_sheet(
                title=nome_servidor,
            )

            planilha.sheet_view.showGridLines = False

            self._preencher_planilha_servidor(
                planilha=planilha,
                naps=naps_por_servidor.get(
                    nome_servidor,
                    [],
                ),
                indice_tabela=indice,
            )

        workbook.save(caminho)

        return caminho

    def _criar_aba_resumo(
        self,
        workbook,
        naps_por_servidor,
        nomes_servidores,
        data_hora,
    ):
        planilha = workbook.create_sheet(
            title="RESUMO",
        )

        planilha.sheet_view.showGridLines = False

        planilha["A1"] = "MGW Monitor"
        planilha["A2"] = "Resumo geral por servidor"
        planilha["A3"] = (
            "Data da coleta: "
            f"{data_hora.strftime('%d/%m/%Y %H:%M:%S')}"
        )

        planilha.merge_cells("A1:F1")
        planilha.merge_cells("A2:F2")
        planilha.merge_cells("A3:F3")

        planilha["A1"].font = Font(
            size=18,
            bold=True,
            color="FFFFFF",
        )

        planilha["A1"].fill = PatternFill(
            fill_type="solid",
            fgColor="1B5E20",
        )

        planilha["A1"].alignment = Alignment(
            horizontal="center",
            vertical="center",
        )

        planilha["A2"].font = Font(
            size=12,
            bold=True,
            color="FFFFFF",
        )

        planilha["A2"].fill = PatternFill(
            fill_type="solid",
            fgColor="2E7D32",
        )

        planilha["A2"].alignment = Alignment(
            horizontal="center",
            vertical="center",
        )

        planilha["A3"].font = Font(
            italic=True,
        )

        planilha["A3"].alignment = Alignment(
            horizontal="center",
            vertical="center",
        )

        cabecalhos = [
            "Servidor",
            "Total",
            "Ativos",
            "Inativos",
            "Com Busy",
            "Com Fail",
        ]

        # Linha 4 vazia.
        planilha.append([])

        # Cabeçalho na linha 5.
        planilha.append(cabecalhos)

        preenchimento_cabecalho = PatternFill(
            fill_type="solid",
            fgColor="2E7D32",
        )

        fonte_cabecalho = Font(
            color="FFFFFF",
            bold=True,
        )

        borda_fina = Side(
            style="thin",
            color="D9D9D9",
        )

        borda = Border(
            left=borda_fina,
            right=borda_fina,
            top=borda_fina,
            bottom=borda_fina,
        )

        for celula in planilha[5]:
            celula.fill = preenchimento_cabecalho
            celula.font = fonte_cabecalho
            celula.border = borda
            celula.alignment = Alignment(
                horizontal="center",
                vertical="center",
            )

        for nome_servidor in nomes_servidores:
            naps_servidor = naps_por_servidor.get(
                nome_servidor,
                [],
            )

            total = len(naps_servidor)

            ativos = sum(
                1
                for nap in naps_servidor
                if nap.esta_ativo()
            )

            inativos = sum(
                1
                for nap in naps_servidor
                if nap.esta_inativo()
            )

            com_busy = sum(
                1
                for nap in naps_servidor
                if nap.busy > 0
            )

            com_fail = sum(
                1
                for nap in naps_servidor
                if nap.fail > 0
            )

            planilha.append(
                [
                    nome_servidor,
                    total,
                    ativos,
                    inativos,
                    com_busy,
                    com_fail,
                ]
            )

            linha_atual = planilha.max_row

            # Destaca em amarelo os servidores
            # que possuem um ou mais NAPs inativos.
            if inativos > 0:
                celula_inativos = planilha.cell(
                    row=linha_atual,
                    column=4,
                )

                celula_inativos.fill = PatternFill(
                    fill_type="solid",
                    fgColor="FFD966",
                )

                celula_inativos.font = Font(
                    bold=True,
                    color="9C5700",
                )

        ultima_linha_dados = planilha.max_row

        planilha.append(
            [
                "TOTAL",
                f"=SUM(B6:B{ultima_linha_dados})",
                f"=SUM(C6:C{ultima_linha_dados})",
                f"=SUM(D6:D{ultima_linha_dados})",
                f"=SUM(E6:E{ultima_linha_dados})",
                f"=SUM(F6:F{ultima_linha_dados})",
            ]
        )

        linha_total = planilha.max_row

        for celula in planilha[linha_total]:
            celula.font = Font(
                bold=True,
            )

            celula.fill = PatternFill(
                fill_type="solid",
                fgColor="C8E6C9",
            )

            celula.border = borda

        for linha in planilha.iter_rows(
            min_row=6,
            max_row=linha_total,
            min_col=1,
            max_col=6,
        ):
            for celula in linha:
                celula.border = borda
                celula.alignment = Alignment(
                    horizontal="center",
                    vertical="center",
                )

        planilha.freeze_panes = "A6"

        planilha.auto_filter.ref = (
            f"A5:F{linha_total}"
        )

        larguras = {
            "A": 18,
            "B": 12,
            "C": 12,
            "D": 14,
            "E": 14,
            "F": 14,
        }

        for coluna, largura in larguras.items():
            planilha.column_dimensions[coluna].width = largura

        planilha.row_dimensions[1].height = 28
        planilha.row_dimensions[2].height = 22
        planilha.row_dimensions[5].height = 24

    def _criar_aba_links_consolidado(
        self,
        workbook,
        naps,
    ):
        planilha = workbook.create_sheet(
            title="Links Consolidado",
        )

        planilha.sheet_view.showGridLines = False

        cabecalhos = [
            "Endereço",
            "Rotas",
            "Servidores afetados",
            "Situação",
        ]

        planilha.append(cabecalhos)

        # A consolidação agora é realizada pelo endereço.
        #
        # Cada endereço aparecerá uma única vez.
        # Todas as rotas associadas a esse endereço
        # serão reunidas dentro da mesma célula.
        enderecos = {}

        for nap in naps:
            endereco = str(
                nap.endereco() or "Não informado"
            ).strip()

            rota = str(
                nap.nome or "Rota não informada"
            ).strip()

            chave_endereco = endereco.casefold()

            if chave_endereco not in enderecos:
                enderecos[chave_endereco] = {
                    "endereco": endereco,
                    "rotas": set(),
                    "servidores": set(),
                    "servidores_inativos": set(),
                }

            # O uso de set elimina rotas repetidas
            # no mesmo endereço.
            enderecos[chave_endereco]["rotas"].add(
                rota
            )

            enderecos[chave_endereco][
                "servidores"
            ].add(nap.servidor)

            if nap.esta_inativo():
                enderecos[chave_endereco][
                    "servidores_inativos"
                ].add(nap.servidor)

        def chave_ordenacao(item):
            endereco = item["endereco"]

            if endereco == "Não informado":
                return (
                    1,
                    endereco,
                )

            return (
                0,
                endereco,
            )

        enderecos_ordenados = sorted(
            enderecos.values(),
            key=chave_ordenacao,
        )

        for item in enderecos_ordenados:
            rotas = sorted(
                item["rotas"],
                key=str.casefold,
            )

            servidores_inativos = sorted(
                item["servidores_inativos"]
            )

            quantidade_inativos = len(
                servidores_inativos
            )

            if quantidade_inativos == 0:
                servidores_afetados = "-"
                situacao = "Funcionando"

            elif quantidade_inativos == 1:
                servidores_afetados = ", ".join(
                    servidores_inativos
                )

                situacao = "Inativo em 1 servidor"

            else:
                servidores_afetados = ", ".join(
                    servidores_inativos
                )

                situacao = (
                    "Inativo em "
                    f"{quantidade_inativos} servidores"
                )

            texto_rotas = "\n".join(rotas)

            planilha.append(
                [
                    item["endereco"],
                    texto_rotas,
                    servidores_afetados,
                    situacao,
                ]
            )

        preenchimento_cabecalho = PatternFill(
            fill_type="solid",
            fgColor="1565C0",
        )

        fonte_cabecalho = Font(
            color="FFFFFF",
            bold=True,
        )

        preenchimento_funcionando = PatternFill(
            fill_type="solid",
            fgColor="D9EAD3",
        )

        preenchimento_inativo = PatternFill(
            fill_type="solid",
            fgColor="F4CCCC",
        )

        borda_fina = Side(
            style="thin",
            color="D9D9D9",
        )

        borda = Border(
            left=borda_fina,
            right=borda_fina,
            top=borda_fina,
            bottom=borda_fina,
        )

        for celula in planilha[1]:
            celula.fill = preenchimento_cabecalho
            celula.font = fonte_cabecalho
            celula.border = borda
            celula.alignment = Alignment(
                horizontal="center",
                vertical="center",
            )

        for linha in planilha.iter_rows(
            min_row=2,
        ):
            situacao = linha[3].value

            if situacao == "Funcionando":
                preenchimento = preenchimento_funcionando
            else:
                preenchimento = preenchimento_inativo

            for celula in linha:
                celula.fill = preenchimento
                celula.border = borda
                celula.alignment = Alignment(
                    vertical="center",
                )

            # Endereço.
            linha[0].alignment = Alignment(
                vertical="center",
            )

            # Rotas, uma por linha dentro da célula.
            linha[1].alignment = Alignment(
                vertical="center",
                wrap_text=True,
            )

            # Servidores afetados.
            linha[2].alignment = Alignment(
                vertical="center",
                wrap_text=True,
            )

            # Situação.
            linha[3].alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True,
            )

            quantidade_rotas = len(
                str(linha[1].value).splitlines()
            )

            quantidade_linhas_servidores = len(
                str(linha[2].value).splitlines()
            )

            maior_quantidade_linhas = max(
                quantidade_rotas,
                quantidade_linhas_servidores,
                1,
            )

            planilha.row_dimensions[
                linha[0].row
            ].height = max(
                20,
                maior_quantidade_linhas * 18,
            )

        planilha.freeze_panes = "A2"
        planilha.auto_filter.ref = planilha.dimensions
        planilha.row_dimensions[1].height = 24

        larguras = {
            "A": 30,
            "B": 55,
            "C": 40,
            "D": 26,
        }

        for coluna, largura in larguras.items():
            planilha.column_dimensions[coluna].width = largura

        if planilha.max_row >= 2:
            tabela = Table(
                displayName="TabelaLinksConsolidado",
                ref=f"A1:D{planilha.max_row}",
            )

            estilo = TableStyleInfo(
                name="TableStyleMedium2",
                showFirstColumn=False,
                showLastColumn=False,
                showRowStripes=False,
                showColumnStripes=False,
            )

            tabela.tableStyleInfo = estilo
            planilha.add_table(tabela)

    def _criar_aba_inativos(
        self,
        workbook,
        naps,
    ):
        planilha = workbook.create_sheet(
            title="NAPs INATIVOS",
        )

        planilha.sheet_view.showGridLines = False

        cabecalhos = [
            "Servidor",
            "ID",
            "Nome",
            "Endereço",
            "Busy",
            "Idle",
            "Fail",
        ]

        planilha.append(cabecalhos)

        inativos = [
            nap
            for nap in naps
            if nap.esta_inativo()
        ]

        inativos_ordenados = sorted(
            inativos,
            key=lambda nap: (
                str(nap.servidor).casefold(),
                str(nap.nome).casefold(),
            ),
        )

        for nap in inativos_ordenados:
            planilha.append(
                [
                    nap.servidor,
                    nap.id,
                    nap.nome,
                    nap.endereco(),
                    nap.busy,
                    nap.idle,
                    nap.fail,
                ]
            )

        preenchimento_cabecalho = PatternFill(
            fill_type="solid",
            fgColor="B71C1C",
        )

        fonte_cabecalho = Font(
            color="FFFFFF",
            bold=True,
        )

        preenchimento_inativo = PatternFill(
            fill_type="solid",
            fgColor="F4CCCC",
        )

        borda_fina = Side(
            style="thin",
            color="D9D9D9",
        )

        borda = Border(
            left=borda_fina,
            right=borda_fina,
            top=borda_fina,
            bottom=borda_fina,
        )

        for celula in planilha[1]:
            celula.fill = preenchimento_cabecalho
            celula.font = fonte_cabecalho
            celula.border = borda
            celula.alignment = Alignment(
                horizontal="center",
                vertical="center",
            )

        for linha in planilha.iter_rows(
            min_row=2,
        ):
            for celula in linha:
                celula.fill = preenchimento_inativo
                celula.border = borda
                celula.alignment = Alignment(
                    vertical="center",
                )

            for indice_coluna in [
                0,
                1,
                4,
                5,
                6,
            ]:
                linha[
                    indice_coluna
                ].alignment = Alignment(
                    horizontal="center",
                    vertical="center",
                )

        planilha.freeze_panes = "A2"
        planilha.auto_filter.ref = planilha.dimensions
        planilha.row_dimensions[1].height = 24

        larguras = {
            "A": 15,
            "B": 10,
            "C": 48,
            "D": 28,
            "E": 12,
            "F": 12,
            "G": 12,
        }

        for coluna, largura in larguras.items():
            planilha.column_dimensions[coluna].width = largura

        if planilha.max_row >= 2:
            tabela = Table(
                displayName="TabelaNAPsInativos",
                ref=f"A1:G{planilha.max_row}",
            )

            estilo = TableStyleInfo(
                name="TableStyleMedium3",
                showFirstColumn=False,
                showLastColumn=False,
                showRowStripes=False,
                showColumnStripes=False,
            )

            tabela.tableStyleInfo = estilo
            planilha.add_table(tabela)

    def _preencher_planilha_servidor(
        self,
        planilha,
        naps,
        indice_tabela,
    ):
        cabecalhos = [
            "ID",
            "Nome",
            "Status",
            "Tipo",
            "Endereço",
            "Busy",
            "Idle",
            "Fail",
        ]

        planilha.append(cabecalhos)

        naps_ordenados = sorted(
            naps,
            key=lambda nap: str(
                nap.nome
            ).casefold(),
        )

        for nap in naps_ordenados:
            status = (
                "Ativo"
                if nap.esta_ativo()
                else "Inativo"
            )

            planilha.append(
                [
                    nap.id,
                    nap.nome,
                    status,
                    self._obter_tipo(
                        nap.tipo
                    ),
                    nap.endereco(),
                    nap.busy,
                    nap.idle,
                    nap.fail,
                ]
            )

        self._formatar_planilha_servidor(
            planilha=planilha,
            indice_tabela=indice_tabela,
        )

    @staticmethod
    def _obter_tipo(tipo):
        tipos = {
            0: "SIP",
            5: "GROUP",
        }

        return tipos.get(
            tipo,
            f"Tipo {tipo}",
        )

    @staticmethod
    def _formatar_planilha_servidor(
        planilha,
        indice_tabela,
    ):
        preenchimento_cabecalho = PatternFill(
            fill_type="solid",
            fgColor="2E7D32",
        )

        fonte_cabecalho = Font(
            color="FFFFFF",
            bold=True,
        )

        preenchimento_ativo = PatternFill(
            fill_type="solid",
            fgColor="D9EAD3",
        )

        preenchimento_inativo = PatternFill(
            fill_type="solid",
            fgColor="F4CCCC",
        )

        borda_fina = Side(
            style="thin",
            color="D9D9D9",
        )

        borda = Border(
            left=borda_fina,
            right=borda_fina,
            top=borda_fina,
            bottom=borda_fina,
        )

        for celula in planilha[1]:
            celula.fill = preenchimento_cabecalho
            celula.font = fonte_cabecalho
            celula.alignment = Alignment(
                horizontal="center",
                vertical="center",
            )
            celula.border = borda

        for linha in planilha.iter_rows(
            min_row=2,
        ):
            status = linha[2].value

            preenchimento = (
                preenchimento_inativo
                if status == "Inativo"
                else preenchimento_ativo
            )

            for celula in linha:
                celula.fill = preenchimento
                celula.border = borda
                celula.alignment = Alignment(
                    vertical="center",
                )

            for indice_coluna in [
                0,
                2,
                3,
                5,
                6,
                7,
            ]:
                linha[
                    indice_coluna
                ].alignment = Alignment(
                    horizontal="center",
                    vertical="center",
                )

        planilha.freeze_panes = "A2"
        planilha.auto_filter.ref = planilha.dimensions
        planilha.row_dimensions[1].height = 24

        larguras_minimas = {
            "A": 10,
            "B": 48,
            "C": 14,
            "D": 14,
            "E": 28,
            "F": 12,
            "G": 12,
            "H": 12,
        }

        for coluna in planilha.columns:
            maior_tamanho = 0

            letra_coluna = get_column_letter(
                coluna[0].column
            )

            for celula in coluna:
                if celula.value is None:
                    continue

                maior_tamanho = max(
                    maior_tamanho,
                    len(str(celula.value)),
                )

            largura_calculada = maior_tamanho + 4

            largura_minima = larguras_minimas.get(
                letra_coluna,
                10,
            )

            planilha.column_dimensions[
                letra_coluna
            ].width = min(
                max(
                    largura_calculada,
                    largura_minima,
                ),
                60,
            )

        if planilha.max_row >= 2:
            tabela = Table(
                displayName=(
                    f"TabelaNAPs{indice_tabela}"
                ),
                ref=f"A1:H{planilha.max_row}",
            )

            estilo = TableStyleInfo(
                name="TableStyleMedium4",
                showFirstColumn=False,
                showLastColumn=False,
                showRowStripes=False,
                showColumnStripes=False,
            )

            tabela.tableStyleInfo = estilo
            planilha.add_table(tabela)