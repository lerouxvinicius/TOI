"""Exercise the full Linux report pipeline with one synthetic invalid number."""
from datetime import datetime, timedelta
from io import BytesIO
from pathlib import Path
import sys
import tempfile

from openpyxl import load_workbook
import xlwt

root = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(root / "TOI_v1.10.6_DUAL_LANE_ML_SHADOW"))
from toi_ipbx_blocked_calls_v1 import build_blocked_calls_report

book = xlwt.Workbook()
sheet = book.add_sheet("IPBX")
headers = [f"Coluna_{n}" for n in range(12)]
headers[9:12] = ["DDD", "Telefone", "Fila"]
for col, value in enumerate(headers):
    sheet.write(0, col, value)
sheet.write(1, 9, "00")
sheet.write(1, 10, "12345")
sheet.write(1, 11, "0247400046")
buffer = BytesIO()
book.save(buffer)
now = datetime.now()
with tempfile.TemporaryDirectory(prefix="report-smoke-", dir=root / "relatorios") as out:
    result = build_blocked_calls_report("0247400046", source=buffer.getvalue(),
                                       start_at=now - timedelta(minutes=2), end_at=now,
                                       output_dir=out, job_id="smoke-synthetic")
    assert result["rows"] == 1
    assert result["carriers"] == [{"name": "N/A", "count": 1, "percent": 100.0}]
    validated = load_workbook(result["output"], read_only=True, data_only=True)
    try:
        sheet = validated.active
        assert sheet["I9"].value == "N/A"
        assert sheet["J9"].value == 1
        assert sheet["K9"].value == 1
    finally:
        validated.close()
print("Linux report, validator, and operator summary OK")
