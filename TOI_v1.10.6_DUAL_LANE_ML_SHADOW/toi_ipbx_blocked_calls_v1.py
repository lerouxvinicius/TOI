"""Busca chamadas por Status IPBX e prepara a consulta de operadoras."""
from __future__ import annotations

from datetime import datetime, timedelta, timezone
from collections import Counter
from html import unescape
from io import BytesIO
import os
from pathlib import Path
import re
import subprocess
import sys
import tempfile
import time
import unicodedata
from urllib.parse import urlsplit, urlunsplit

from bs4 import BeautifulSoup
from dotenv import load_dotenv
from playwright.sync_api import TimeoutError as PlaywrightTimeoutError, sync_playwright
import xlrd
import xlwt
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo

from toi_ipbx_live_trace_v1 import OLE_SIGNATURE, log_trace

ROOT = Path(__file__).resolve().parent.parent
OUTPUT_DIR = Path(os.getenv("TOI_BLOCKED_CALLS_DIR", str(ROOT / "relatorios" / "consulta_telefone")))
QUEUES = {
    "0247400046": "DISC_LOCALIZADOR_VIRTUAL_05_2026",
    "0247400047": "DISC_LOCALIZADOR_VIRTUAL_05_2026_2",
}
PROJECT = "1474"
DEFAULT_STATUS = "21"
IPBX_STATUSES = {"21": "Rel21 - Call Rejected", "16": "Rel16 - Normal Call Clearing",
                 "34": "Rel34 - No Circuit Available"}


def validate_ipbx_status(value) -> str:
    status = str(value)
    if status not in IPBX_STATUSES:
        raise ValueError("Selecione Rel21, Rel16 ou Rel34 - No Circuit Available")
    return status


def _verify_period_controls(page, start_at: datetime, end_at: datetime) -> None:
    """Wait for the Portal's hidden timestamps, which update after the visible fields."""
    checks = []
    for name, stamp in (("cldDtStart", start_at), ("txtDtStart", start_at),
                        ("cldDtFinish", end_at), ("txtDtFinish", end_at)):
        if name.startswith("cld"):
            visible = stamp.strftime("%d/%m/%Y")
            expected = stamp.replace(hour=0, minute=0, second=0, microsecond=0)
        else:
            visible = stamp.strftime("%H:%M:%S")
            expected = stamp.replace(microsecond=0)
        milliseconds = str(int(expected.replace(tzinfo=timezone.utc).timestamp() * 1000))
        checks.append((name, visible, milliseconds))
    predicate = """checks => checks.every(([name, visible, stamp]) => {
        const input = document.getElementsByName('ctl09$' + name)[0];
        const hidden = document.getElementsByName('ctl09_' + name + '_TS')[0];
        return input && hidden && input.value === visible && hidden.value === stamp;
    })"""
    for attempt in range(2):
        try:
            page.wait_for_function(predicate, arg=checks, timeout=5_000 if attempt == 0 else 15_000)
            return
        except PlaywrightTimeoutError as exc:
            if attempt:
                raise RuntimeError("O Portal não confirmou o período selecionado para Rel34") from exc
            for name, visible, _ in checks:
                control = page.locator(f'[name="ctl09${name}"]')
                control.fill(visible)
                control.press("Tab")


def _portal_report_url(portal, home) -> str:
    soup = BeautifulSoup(home.text, "html.parser")
    anchor = next(
        (a for a in soup.find_all("a")
         if a.get("href", "").startswith("javascript:openReport")
         and "pIdMenu=8077" in a.get("href", "")), None,
    )
    if anchor is None:
        raise RuntimeError("Relatório Ligações e Gravações - IPBX indisponível")
    match = re.search(r'openReport\s*\(\s*["\'][^"\']+["\']\s*,\s*["\']([^"\']+)',
                      anchor["href"])
    if match is None:
        raise RuntimeError("Link do relatório IPBX não reconhecido")
    parsed = urlsplit(home.url)
    path = "/" + unescape(match.group(1)).replace("\\", "/").lstrip("/")
    return urlunsplit((parsed.scheme, parsed.netloc, path, "", ""))


def fetch_blocked_calls(queue: str, start_at: datetime, end_at: datetime,
                        job_id: str = "manual", *, ipbx_status=DEFAULT_STATUS) -> tuple[bytes, str]:
    """Exporta a fila escolhida no período informado pelo usuário."""
    if queue not in QUEUES:
        raise ValueError("Fila IPBX inválida")
    ipbx_status = validate_ipbx_status(ipbx_status)
    load_dotenv(ROOT / ".env")
    user, password = os.getenv("ORBE_LOGIN"), os.getenv("ORBE_PASSWORD")
    if not user or not password:
        raise RuntimeError("Credenciais ORBE_LOGIN/ORBE_PASSWORD ausentes")
    from monitor_portal_orbe_v1_gerencial import PortalIPBXCanaisTravados

    portal = PortalIPBXCanaisTravados(
        usuario=user, senha=password,
        config_instancias={"calls": {"projeto": "869", "menu": "8077"}},
    )
    log_trace("portal_login_start", queue, job_id)
    portal.autenticar()
    _, home = portal._selecionar_projeto("calls")
    report_url = _portal_report_url(portal, home)
    with sync_playwright() as playwright:
        browser = playwright.chromium.launch(headless=True)
        try:
            context = browser.new_context(accept_downloads=True)
            context.add_cookies([
                {"name": c.name, "value": c.value, "domain": c.domain,
                 "path": c.path or "/", "secure": c.secure, "httpOnly": False}
                for c in portal.session.cookies
            ])
            page = context.new_page()
            page.goto(report_url, wait_until="domcontentloaded", timeout=30_000)
            if "DynamicQuery.aspx" not in page.url:
                raise RuntimeError("O Portal não abriu o formulário de Ligações e Gravações")
            page.locator('[name="ctl12$comboBox"]').select_option(PROJECT)
            page.wait_for_function(
                'document.querySelectorAll("[name=\'ctl15$Lst_lstCodeQueue\'] option").length > 0',
                timeout=30_000,
            )
            for field, stamp in (("cldDtStart", start_at), ("cldDtFinish", end_at)):
                control = page.locator(f'[name="ctl09${field}"]')
                control.fill(stamp.strftime("%d/%m/%Y"))
                control.press("Tab")
            for field, stamp in (("txtDtStart", start_at), ("txtDtFinish", end_at)):
                control = page.locator(f'[name="ctl09${field}"]')
                control.fill(stamp.strftime("%H:%M:%S"))
                control.press("Tab")
            page.locator('[name="ctl15$Chk_lstCodeQueue"]').check()
            page.locator('[name="ctl15$Lst_lstCodeQueue"]').select_option(queue)
            page.locator('[name="ctl33$Chk_lstCodeStatusIpbx"]').check()
            status_control = page.locator('[name="ctl33$Lst_lstCodeStatusIpbx"]')
            status_control.select_option(ipbx_status)
            if status_control.input_value() != ipbx_status:
                raise RuntimeError("O Portal não selecionou o Status IPBX solicitado")
            page.locator("#cmbIdFormatTypeTop_B-1").click()
            page.locator("#cmbIdFormatTypeTop_DDD_L_LBI2T").click()
            if page.locator('[name="cmbIdFormatTypeTop"]').input_value() != "XLS":
                raise RuntimeError("O Portal não selecionou o formato XLS")
            if ipbx_status == "34":
                _verify_period_controls(page, start_at, end_at)
            log_trace("report_query_start", queue, job_id,
                      start=start_at.isoformat(), end=end_at.isoformat(), ipbx_status=ipbx_status)
            with page.expect_download(timeout=180_000) as transfer:
                page.locator('[name="btnQueryTop"]').click(timeout=180_000)
            download = transfer.value
            contents = Path(download.path()).read_bytes()
            if not contents.startswith(OLE_SIGNATURE):
                raise RuntimeError("O Portal não entregou uma planilha XLS válida")
            name = Path(download.suggested_filename).name
            if not re.fullmatch(r"869-Ligações[eE]Gravações-IPBX-\d{14}\.xls", name):
                name = "869-LigaçõeseGravações-IPBX-" + datetime.now().strftime("%d%m%Y%H%M%S") + ".xls"
            log_trace("report_export_ready", queue, job_id, bytes=len(contents))
            return contents, name
        finally:
            browser.close()


def _digits(value) -> str:
    if value is None or value == "":
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return re.sub(r"\D", "", str(value))


def add_phone_column(source: bytes, queue: str) -> tuple[bytes, int]:
    """Insere L=DDD+Telefone, renomeando K para Tel sem perder as demais colunas."""
    book = xlrd.open_workbook(file_contents=source)
    if not book.nsheets:
        raise ValueError("Planilha IPBX vazia")
    writer = xlwt.Workbook(encoding="utf-8")
    date_style = xlwt.easyxf(num_format_str="DD/MM/YYYY HH:MM:SS.000")
    phone_style = xlwt.easyxf(num_format_str="@")
    queue_style = xlwt.easyxf(num_format_str="0000000000")
    total = 0
    for sheet in book.sheets():
        if sheet.nrows == 0 or sheet.ncols < 12:
            raise ValueError("Layout IPBX incompleto")
        if str(sheet.cell_value(0, 9)).strip().upper() != "DDD" or \
                str(sheet.cell_value(0, 10)).strip().lower() != "telefone":
            raise ValueError("Layout IPBX inválido: esperadas DDD em J e Telefone em K")
        out = writer.add_sheet(sheet.name[:31])
        for row in range(sheet.nrows):
            for col in range(sheet.ncols):
                target = col + 1 if col >= 11 else col
                cell = sheet.cell(row, col)
                value = cell.value
                style = date_style if cell.ctype == xlrd.XL_CELL_DATE else xlwt.Style.default_style
                if row == 0 and col == 10:
                    value = "Tel"
                if row > 0 and col == 11:
                    style = queue_style
                    if _digits(value).zfill(10) != queue:
                        raise ValueError("A planilha contém chamadas de outra fila")
                out.write(row, target, value, style)
            if row == 0:
                out.write(row, 11, "Telefone")
            else:
                ddd = _digits(sheet.cell_value(row, 9))
                tel = _digits(sheet.cell_value(row, 10))
                out.write(row, 11, ddd + tel if ddd and tel else "", phone_style)
                total += 1
        out.col(10).width = 4200
        out.col(11).width = 5200
    data = BytesIO()
    writer.save(data)
    return data.getvalue(), total


def _carrier_summary(path: Path, expected_rows: int) -> list[dict]:
    """Resume a mesma coluna F usada pela Tabela Dinâmica do Excel."""
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = book.worksheets[0]
        if sheet.cell(1, 6).value != "Operadora_Atual":
            raise RuntimeError("A operadora não está na coluna F da planilha validada")
        counts = Counter(str(row[0]).strip() if row[0] is not None else "(vazio)"
                         for row in sheet.iter_rows(min_row=2, min_col=6, max_col=6,
                                                    values_only=True))
        if sum(counts.values()) != expected_rows:
            raise RuntimeError("A quantidade validada difere da planilha IPBX")
        return [{"name": name, "count": count, "percent": 100 * count / expected_rows}
                for name, count in sorted(counts.items(), key=lambda item: (-item[1], item[0]))]
    finally:
        book.close()


def _normalized_heading(value) -> str:
    value = unicodedata.normalize("NFKD", str(value or ""))
    return re.sub(r"[^a-z0-9]", "", "".join(c for c in value if not unicodedata.combining(c)).lower())


def _fetch_classification(queue: str, start_at: datetime, end_at: datetime,
                          job_id: str, *, historical: bool) -> bytes:
    """Export the realtime or historical classification report for Rel34."""
    if queue not in QUEUES:
        raise ValueError("Fila IPBX inválida")
    load_dotenv(ROOT / ".env")
    user, password = os.getenv("ORBE_LOGIN"), os.getenv("ORBE_PASSWORD")
    if not user or not password:
        raise RuntimeError("Credenciais ORBE_LOGIN/ORBE_PASSWORD ausentes")
    from monitor_portal_orbe_v1_gerencial import PortalIPBXCanaisTravados

    menu = "7960" if historical else "7977"
    portal = PortalIPBXCanaisTravados(
        usuario=user, senha=password,
        config_instancias={"classification": {"projeto": "868", "menu": menu}},
    )
    portal.autenticar()
    _, home = portal._selecionar_projeto("classification")
    soup = BeautifulSoup(home.text, "html.parser")
    anchor = next((a for a in soup.find_all("a")
                   if a.get("href", "").startswith("javascript:openReport")
                   and f"pIdMenu={menu}&" in a.get("href", "")
                   and len(a.get_text(" ", strip=True)) < 100), None)
    if anchor is None:
        raise RuntimeError("Relatório Telefonia - Classificação por Status indisponível")
    match = re.search(r'openReport\s*\(\s*["\'][^"\']+["\']\s*,\s*["\']([^"\']+)', anchor["href"])
    if match is None:
        raise RuntimeError("Link do relatório de classificação não reconhecido")
    parsed = urlsplit(home.url)
    path = "/" + unescape(match.group(1)).replace("\\", "/").lstrip("/")
    report_url = urlunsplit((parsed.scheme, parsed.netloc, path, "", ""))
    with sync_playwright() as playwright:
        browser = playwright.chromium.launch(headless=True)
        try:
            context = browser.new_context(accept_downloads=True)
            context.add_cookies([
                {"name": c.name, "value": c.value, "domain": c.domain,
                 "path": c.path or "/", "secure": c.secure, "httpOnly": False}
                for c in portal.session.cookies
            ])
            page = context.new_page()
            page.goto(report_url, wait_until="domcontentloaded", timeout=30_000)
            if "DynamicQuery.aspx" not in page.url:
                raise RuntimeError("O Portal não abriu o relatório de classificação")
            for field, stamp in (("cldDtStart", start_at), ("cldDtFinish", end_at)):
                control = page.locator(f'[name="ctl09${field}"]')
                control.fill(stamp.strftime("%d/%m/%Y"))
                control.press("Tab")
            for field, stamp in (("txtDtStart", start_at), ("txtDtFinish", end_at)):
                control = page.locator(f'[name="ctl09${field}"]')
                control.fill(stamp.strftime("%H:%M:%S"))
                control.press("Tab")
            checked = ("ctl39",) if historical else ("ctl18", "ctl21", "ctl33")
            unchecked = (("ctl12", "ctl15", "ctl18", "ctl21", "ctl24", "ctl27", "ctl30", "ctl33", "ctl36", "ctl42")
                         if historical else ("ctl12", "ctl15", "ctl24", "ctl27", "ctl30"))
            for field in checked:
                page.locator(f'[name="{field}$checkBox"]').check()
            for field in unchecked:
                if page.locator(f'[name="{field}$checkBox"]').is_checked():
                    raise RuntimeError("Agrupamento inesperado no relatório de classificação")
            if historical:
                page.locator('[name="ctl48$comboBox"]').select_option(PROJECT)
                queue_group = "ctl51"
            else:
                page.locator('[name="ctl39$Chk_cmbIdProject"]').check()
                page.locator('[name="ctl39$Lst_cmbIdProject"]').select_option(PROJECT)
                queue_group = "ctl42"
            page.wait_for_function(
                f'document.querySelectorAll("[name=\'{queue_group}$Lst_lstCodeQueue\'] option").length > 0',
                timeout=30_000,
            )
            page.locator(f'[name="{queue_group}$Chk_lstCodeQueue"]').check()
            page.locator(f'[name="{queue_group}$Lst_lstCodeQueue"]').select_option(queue)
            if page.locator(f'[name="{queue_group}$Lst_lstCodeQueue"]').input_value() != queue:
                raise RuntimeError("O Portal não selecionou a fila no relatório de classificação")
            page.locator("#cmbIdFormatTypeTop_B-1").click()
            page.locator("#cmbIdFormatTypeTop_DDD_L_LBI2T").click()
            if page.locator('[name="cmbIdFormatTypeTop"]').input_value() != "XLS":
                raise RuntimeError("O Portal não selecionou XLS no relatório de classificação")
            _verify_period_controls(page, start_at, end_at)
            kind = "historical" if historical else "realtime"
            log_trace("classification_query_start", queue, job_id, source=kind,
                      start=start_at.isoformat(), end=end_at.isoformat(), ipbx_status="34")
            with page.expect_download(timeout=180_000) as transfer:
                page.locator('[name="btnQueryTop"]').click(no_wait_after=True)
            contents = Path(transfer.value.path()).read_bytes()
            if not contents.startswith(OLE_SIGNATURE):
                raise RuntimeError("O Portal não entregou uma planilha de classificação XLS válida")
            log_trace("classification_export_ready", queue, job_id, source=kind, bytes=len(contents))
            return contents
        finally:
            browser.close()


def fetch_realtime_classification(queue: str, start_at: datetime, end_at: datetime,
                                  job_id: str = "manual") -> bytes:
    return _fetch_classification(queue, start_at, end_at, job_id, historical=False)


def fetch_historical_classification(queue: str, start_at: datetime, end_at: datetime,
                                    job_id: str = "manual") -> bytes:
    return _fetch_classification(queue, start_at, end_at, job_id, historical=True)


def _historical_route_counts(source: bytes) -> Counter:
    """Use A=Rota, B=Status IPBX and E=Ligações in the historical export."""
    book = xlrd.open_workbook(file_contents=source)
    if book.nsheets != 1:
        raise ValueError("A classificação histórica deve conter uma única aba")
    sheet = book.sheet_by_index(0)
    if sheet.ncols < 5 or _normalized_heading(sheet.cell_value(0, 0)) != "rota" or \
            _normalized_heading(sheet.cell_value(0, 1)) != "statusipbx" or \
            _normalized_heading(sheet.cell_value(0, 4)) != "ligacoes":
        raise ValueError("Layout histórico inválido: esperadas Rota em A, Status IPBX em B e Ligações em E")
    routes = Counter()
    for row in range(1, sheet.nrows):
        status = str(sheet.cell_value(row, 1)).strip()
        if not re.match(r"^34\s*-", status):
            continue
        amount = sheet.cell_value(row, 4)
        if not isinstance(amount, (int, float)) or amount < 0 or int(amount) != amount:
            raise ValueError("Quantidade inválida no relatório histórico")
        route = str(sheet.cell_value(row, 0)).strip() or "(vazio)"
        routes[route] += int(amount)
    return routes


def _ipbx_dimension_counts(source: bytes) -> dict[str, Counter]:
    """Count trunk groups from X=Operadora and gateways from Y=Gateway."""
    book = xlrd.open_workbook(file_contents=source)
    if book.nsheets != 1:
        raise ValueError("A exportação IPBX deve conter uma única aba")
    sheet = book.sheet_by_index(0)
    if sheet.ncols < 25 or _normalized_heading(sheet.cell_value(0, 23)) != "operadora" or \
            _normalized_heading(sheet.cell_value(0, 24)) != "gateway":
        raise ValueError("Layout IPBX inválido: esperadas Operadora em X e Gateway em Y")
    result = {"trunk": Counter(), "gateway": Counter()}
    for row in range(1, sheet.nrows):
        result["trunk"][str(sheet.cell_value(row, 23)).strip() or "(vazio)"] += 1
        result["gateway"][str(sheet.cell_value(row, 24)).strip() or "(vazio)"] += 1
    return result


def _style_teal_header(sheet, cell_range: str) -> None:
    fill = PatternFill("solid", fgColor="00B1B0")
    font = Font(name="Calibri", size=15, bold=True, color="FFFFFF")
    for row in sheet[cell_range]:
        for cell in row:
            cell.fill = fill
            cell.font = font
            cell.alignment = Alignment(vertical="center")


def _add_dimension_summary(book, carriers, counts, queue, start_at, end_at,
                           classification_source):
    """Build four filterable counts and show their independent source totals."""
    sheet = book.create_sheet("Resumo por Dimensão")
    sheet.sheet_view.showGridLines = False
    sheet.append(["Busca de chamadas · resumo por dimensão"])
    for row, label, value in ((2, "Fila", queue), (3, "Status IPBX", IPBX_STATUSES["34"]),
                              (4, "Data Início", start_at.strftime("%d/%m/%Y %H:%M:%S") if start_at else ""),
                              (5, "Data Fim", end_at.strftime("%d/%m/%Y %H:%M:%S") if end_at else ""),
                              (6, "Operadoras: Ligações e Gravações - IPBX", sum(item["count"] for item in carriers)),
                              (7, "Rotas: " + classification_source, sum(counts["route"].values())),
                              (8, "Grupo de Troncos (X) / Gateway (Y): Ligações e Gravações - IPBX",
                               sum(counts["trunk"].values()))):
        sheet.cell(row, 1, label)
        sheet.cell(row, 2, value)
    sections = ((1, "Operadora", Counter({x["name"]: x["count"] for x in carriers}), "Operadoras"),
                (4, "Rota", counts["route"], "Rotas"),
                (7, "Grupo de Troncos", counts["trunk"], "GruposTroncos"),
                (10, "Gateway", counts["gateway"], "Gateways"))
    for col, title, values, table_name in sections:
        sheet.cell(9, col, title)
        sheet.cell(9, col + 1, "Quantidade")
        for row, (name, count) in enumerate(sorted(values.items(), key=lambda x: (-x[1], x[0])), 10):
            sheet.cell(row, col, name)
            sheet.cell(row, col + 1, count)
        last = 9 + len(values)
        if values:
            table = Table(displayName=table_name, ref=f"{sheet.cell(9,col).column_letter}9:{sheet.cell(last,col+1).column_letter}{last}")
            table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
            sheet.add_table(table)
        sheet.cell(last + 1, col, "Total Geral")
        sheet.cell(last + 1, col + 1, sum(values.values()))
        sheet.cell(last + 1, col).font = Font(bold=True)
        sheet.cell(last + 1, col + 1).font = Font(bold=True)
        sheet.column_dimensions[sheet.cell(9, col).column_letter].width = 34
        sheet.column_dimensions[sheet.cell(9, col + 1).column_letter].width = 15
    for area in ("A1:C1", "A4:C5", "A9:B9", "D9:E9", "G9:H9", "J9:K9"):
        _style_teal_header(sheet, area)
    for row in (1, 4, 5, 9):
        sheet.row_dimensions[row].height = 27
    sheet.freeze_panes = "A10"


def add_operator_summary(path: Path, carriers: list[dict], total: int, *,
                         classification=None, queue=None, ipbx_status=DEFAULT_STATUS,
                         start_at=None, end_at=None,
                         classification_source="Telefonia - Classificação por Status") -> None:
    """Create the same operator counts and percentages on Linux without Excel COM."""
    book = load_workbook(path)
    try:
        sheet = book.worksheets[0]
        sheet.title = "Operadoras"
        sheet.sheet_view.showGridLines = False
        sheet.auto_filter.ref = f"A1:G{sheet.max_row}"
        sheet.freeze_panes = "A2"
        for column, width in {"A": 24, "B": 24, "C": 23, "D": 16, "E": 22,
                              "F": 23, "G": 16, "H": 3, "I": 27,
                              "J": 36 if ipbx_status == "34" else 18, "K": 17}.items():
            sheet.column_dimensions[column].width = width
        _style_teal_header(sheet, "A1:G1")
        sheet.row_dimensions[1].height = 28
        if ipbx_status == "34":
            for row, label, value in ((3,"Fila",queue),(4,"Status IPBX",IPBX_STATUSES["34"]),
                                      (5,"Data Início",start_at.strftime("%d/%m/%Y %H:%M:%S")),
                                      (6,"Data Fim",end_at.strftime("%d/%m/%Y %H:%M:%S"))):
                sheet.cell(row,9,label)
                sheet.cell(row,10,value)
        for column, value in zip(("I8", "J8", "K8"),
                                 ("Operadora_Atual", "Quantidade", "Percentual")):
            sheet[column] = value
        for index, item in enumerate(carriers, 9):
            sheet.cell(index, 9, item["name"])
            sheet.cell(index, 10, item["count"])
            sheet.cell(index, 11, item["count"] / total if total else 0).number_format = "0%"
        last = 9 + len(carriers)
        sheet.cell(last, 9, "Total Geral")
        sheet.cell(last, 10, total)
        sheet.cell(last, 11, 1 if total else 0).number_format = "0%"
        if carriers:
            table = Table(displayName="OperadorasPorTelefone", ref=f"I8:K{last-1}")
            table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
            sheet.add_table(table)
        _style_teal_header(sheet, "I8:K8")
        sheet.row_dimensions[8].height = 27
        for col in (9, 10, 11):
            sheet.cell(last, col).font = Font(name="Calibri", size=11, bold=True)
        if ipbx_status == "34":
            if classification is None:
                raise RuntimeError("Classificação Rel34 não fornecida")
            _add_dimension_summary(book, carriers, classification, queue, start_at, end_at,
                                   classification_source)
        book.save(path)
    finally:
        book.close()


def build_blocked_calls_report(queue: str, *, start_at=None, end_at=None,
                               report_date=None, source=None,
                               historical_source=None,
                               validator=None, output_dir=None, progress=None,
                               job_id: str = "manual", ipbx_status=DEFAULT_STATUS) -> dict:
    if queue not in QUEUES:
        raise ValueError("Fila IPBX inválida")
    ipbx_status = validate_ipbx_status(ipbx_status)
    if start_at is None and end_at is None:
        day = report_date or (datetime.now().date() - timedelta(days=1))
        start_at = datetime.combine(day, datetime.min.time())
        end_at = start_at.replace(hour=23, minute=59, second=59)
    elif not isinstance(start_at, datetime) or not isinstance(end_at, datetime):
        raise ValueError("Informe Data Início e Data Fim")
    if end_at <= start_at:
        raise ValueError("Data Fim deve ser posterior à Data Início")
    day = start_at.date()
    if progress:
        progress("Extraindo chamadas do Portal Orbe · " + IPBX_STATUSES[ipbx_status])
    if source is None:
        source, filename = fetch_blocked_calls(queue, start_at, end_at, job_id,
                                               ipbx_status=ipbx_status)
    else:
        filename = "869-LigaçõeseGravações-IPBX-" + datetime.now().strftime("%d%m%Y%H%M%S") + ".xls"
    if progress:
        progress("Preparando a coluna Telefone")
    prepared, rows = add_phone_column(source, queue)
    classification = None
    classification_source = "Telefonia - Classificação por Status"
    if ipbx_status == "34":
        classification = _ipbx_dimension_counts(source)
        if sum(classification["trunk"].values()) != rows or \
                sum(classification["gateway"].values()) != rows:
            raise RuntimeError("Grupo de Troncos e Gateway não cobrem as chamadas IPBX extraídas")
        if progress:
            progress("Extraindo as Rotas da Telefonia - Classificação por Status")
        if historical_source is None:
            historical_source = fetch_historical_classification(queue, start_at, end_at, job_id)
        classification["route"] = _historical_route_counts(historical_source)
    target = Path(output_dir) if output_dir is not None else OUTPUT_DIR
    target.mkdir(parents=True, exist_ok=True)
    script = Path(validator) if validator is not None else Path(__file__).with_name("validador_telefones.py")
    with tempfile.TemporaryDirectory(prefix="consulta_", dir=target) as tmp:
        work = Path(tmp)
        (work / "telefones.xls").write_bytes(prepared)
        log_trace("carrier_validation_start", queue, job_id, rows=rows)
        if progress:
            progress(f"Consultando operadoras de {rows} chamadas")
        process = subprocess.Popen([sys.executable, str(script)], cwd=work,
                                   stdout=subprocess.PIPE, stderr=subprocess.STDOUT,
                                   text=True, encoding="utf-8", errors="replace",
                                   env={**os.environ, "PYTHONIOENCODING": "utf-8"})
        started = time.monotonic()
        last_progress = -1
        progress_file = work / "consulta_progresso.txt"
        while process.poll() is None:
            if progress and progress_file.exists():
                try:
                    done = int(progress_file.read_text(encoding="ascii"))
                except (ValueError, OSError):
                    done = last_progress
                if done != last_progress:
                    progress(f"Consultando operadoras: {done}/{rows}")
                    last_progress = done
            if time.monotonic() - started > 14400:
                process.kill()
                process.communicate()
                raise RuntimeError("O validador excedeu quatro horas de execução")
            time.sleep(2)
        stdout, _ = process.communicate()
        validated = work / "telefones_validados_consultaoperadora.xlsx"
        if process.returncode or not validated.is_file():
            raise RuntimeError("O validador de telefones falhou: " +
                               (stdout.strip()[-300:] or "planilha não gerada"))
        if progress:
            progress("Calculando percentuais por operadora")
        carriers = _carrier_summary(validated, rows)
        add_operator_summary(validated, carriers, rows, classification=classification,
                             queue=queue, ipbx_status=ipbx_status,
                             start_at=start_at, end_at=end_at,
                             classification_source=classification_source)
        (target / filename).write_bytes(source)
        os.replace(work / "telefones.xls", target / "telefones.xls")
        destination = target / validated.name
        deadline = time.monotonic() + 600
        while True:
            try:
                os.replace(validated, destination)
                break
            except PermissionError as exc:
                if time.monotonic() >= deadline:
                    recovery = target / f"telefones_validados_recuperacao_{job_id[:12]}.xlsx"
                    os.replace(validated, recovery)
                    raise RuntimeError(
                        "A planilha anterior ainda está em uso. "
                        f"O resultado validado foi preservado em {recovery}"
                    ) from exc
                if progress:
                    progress("Aguardando liberação da planilha anterior")
                time.sleep(3)
    log_trace("job_complete", queue, job_id, rows=rows, ipbx_status=ipbx_status)
    result = {"queue": queue, "date": day.isoformat(), "rows": rows,
            "carriers": carriers, "ipbx_status": ipbx_status, "ipbx_status_label": IPBX_STATUSES[ipbx_status],
            "period": start_at.strftime("%d/%m/%Y %H:%M:%S") + " a " +
                      end_at.strftime("%d/%m/%Y %H:%M:%S"),
            "portal_limit_possible": rows >= 1000,
            "source": str(target / filename),
            "input": str(target / "telefones.xls"),
            "output": str(target / "telefones_validados_consultaoperadora.xlsx")}
    if ipbx_status == "34":
        result["classification_source"] = classification_source
    return result
