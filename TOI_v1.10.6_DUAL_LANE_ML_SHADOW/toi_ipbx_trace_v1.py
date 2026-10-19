"""Gera relatório IPBX sob demanda de um Trace de Chamadas, sem rede ou persistência."""
from __future__ import annotations

from collections import Counter
from io import BytesIO, StringIO
import secrets

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from toi_ipbx_operadoras_v1 import STATUS_RE, estimar_operadora, normalizar


MAX_UPLOAD_BYTES = 5 * 1024 * 1024


def _rows(contents):
    if contents.startswith(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"):
        import xlrd
        workbook = xlrd.open_workbook(file_contents=contents, on_demand=True,
                                      logfile=StringIO())
        try:
            sheet = workbook.sheet_by_index(0)
            for index in range(sheet.nrows):
                yield sheet.row_values(index)
        finally:
            workbook.release_resources()
    elif contents.startswith(b"PK\x03\x04"):
        workbook = load_workbook(BytesIO(contents), read_only=True, data_only=True)
        try:
            yield from workbook.worksheets[0].iter_rows(values_only=True)
        finally:
            workbook.close()
    else:
        raise ValueError("Arquivo inválido: selecione o XLS ou XLSX do Trace de Chamadas")


def _cell_text(value):
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value or "")


def analyze(contents, status, *, classifier=estimar_operadora, rng=None):
    status = str(status)
    if status not in ("19", "21"):
        raise ValueError("Status IPBX deve ser 19 ou 21")
    if not contents or len(contents) > MAX_UPLOAD_BYTES:
        raise ValueError("O arquivo deve ter até 5 MB")
    source = iter(_rows(contents))
    try:
        header = next(source)
    except StopIteration:
        raise ValueError("Trace de Chamadas sem cabeçalho") from None
    if len(header) <= 10 or not _cell_text(header[8]).lower().startswith("telefone") or \
            not _cell_text(header[10]).lower().startswith("status ipbx"):
        raise ValueError("Layout do Trace inválido: esperadas Telefone na coluna I e Status IPBX na K")

    matching = invalid = 0
    distinct = {}
    for row in source:
        if len(row) <= 10:
            continue
        match = STATUS_RE.match(_cell_text(row[10]))
        if not match or match.group(1) != status:
            continue
        matching += 1
        phone = normalizar(_cell_text(row[8]))
        if phone:
            distinct[phone] = None
        else:
            invalid += 1
    if not matching:
        raise ValueError(f"O arquivo não contém chamadas com Status IPBX {status}")
    if not distinct:
        raise ValueError(f"Nenhum telefone válido encontrado na coluna I para o Status {status}")
    selection = (rng or secrets.SystemRandom()).sample(list(distinct), min(100, len(distinct)))
    records = []
    for phone in selection:
        carrier, method = classifier(phone)
        records.append({"phone": phone, "ddd": phone[:2], "carrier": carrier,
                        "method": method, "status": status})
    return {"status": status, "matching_rows": matching, "valid_distinct": len(distinct),
            "invalid_rows": invalid, "selected": len(records), "records": records}


def workbook_bytes(result):
    records = result["records"]
    counts = Counter(row["carrier"] for row in records)
    doc = Workbook()
    summary = doc.active
    summary.title = "Resumo"
    summary.append([f"IPBX {result['status']} · Trace de Chamadas"])
    summary.append(["Estimativa local por prefixo; não confirma operadora atual após portabilidade."])
    summary.append([])
    for label, value in (("Linhas do status no arquivo", result["matching_rows"]),
                         ("Telefones distintos válidos", result["valid_distinct"]),
                         ("Telefones selecionados aleatoriamente", result["selected"]),
                         ("Linhas sem telefone válido", result["invalid_rows"])):
        summary.append([label, value])
    summary.append(["Amostra incompleta" if result["selected"] < 100 else "Amostra completa",
                    "Menos de 100 telefones distintos disponíveis" if result["selected"] < 100 else "100 telefones distintos"])
    summary.append([])
    summary.append(["Operadora estimada", "Números", "% dos selecionados"])
    for name in ("GAMA", "BETA", "ALFA", "OUTRA / DESCONHECIDA", "DESCONHECIDA"):
        amount = counts[name]
        summary.append([name, amount, amount / len(records)])
        summary.cell(summary.max_row, 3).number_format = "0.00%"
    summary.column_dimensions["A"].width = 61
    summary.column_dimensions["B"].width = 26
    summary.column_dimensions["C"].width = 22

    detail = doc.create_sheet("Telefones")
    detail.append(["Telefone com DDD", "DDD", "Operadora estimada", "Método local", "Status IPBX"])
    for row in records:
        detail.append([row["phone"], row["ddd"], row["carrier"], row["method"], row["status"]])
        detail.cell(detail.max_row, 1).number_format = "@"
    for column, width in {"A": 24, "B": 12, "C": 27, "D": 37, "E": 16}.items():
        detail.column_dimensions[column].width = width
    detail.freeze_panes = "A2"
    detail.auto_filter.ref = detail.dimensions

    ddd_sheet = doc.create_sheet("DDD")
    ddd_sheet.append(["DDD", "Operadora estimada", "Números", "% dos selecionados"])
    for (ddd, carrier), amount in sorted(Counter((x["ddd"], x["carrier"]) for x in records).items()):
        ddd_sheet.append([ddd, carrier, amount, amount / len(records)])
        ddd_sheet.cell(ddd_sheet.max_row, 4).number_format = "0.00%"
    for column, width in {"A": 12, "B": 28, "C": 16, "D": 23}.items():
        ddd_sheet.column_dimensions[column].width = width

    for sheet in doc.worksheets:
        sheet.sheet_view.showGridLines = False
        for cell in sheet[1]:
            cell.fill = PatternFill("solid", fgColor="0B2636")
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(vertical="center")
        sheet.row_dimensions[1].height = 25
    output = BytesIO()
    doc.save(output)
    return output.getvalue()
