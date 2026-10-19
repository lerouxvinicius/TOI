"""Busca o Trace autenticado no Portal Orbe e consulta a operadora atual sob demanda."""
from __future__ import annotations

from html import unescape
from io import BytesIO
from pathlib import Path
import logging
from logging.handlers import RotatingFileHandler
import os
import random
import re
import sys
import time
from urllib.parse import urljoin

from bs4 import BeautifulSoup
from dotenv import load_dotenv
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from playwright.sync_api import sync_playwright
import requests

from toi_ipbx_trace_v1 import analyze


ROOT = Path(__file__).resolve().parent.parent
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))
from monitor_portal_orbe_v1_gerencial import PortalIPBXCanaisTravados  # noqa: E402

CONSULTA_URL = "https://consulta-operadora.example.com/site2015/resposta.php"
OLE_SIGNATURE = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"
LOG_PATH = ROOT / "logs" / "toi_ipbx_trace.log"
LOG_PATH.parent.mkdir(parents=True, exist_ok=True)
logger = logging.getLogger("toi.ipbx.trace")
if not logger.handlers:
    handler = RotatingFileHandler(LOG_PATH, maxBytes=2_000_000, backupCount=4, encoding="utf-8")
    handler.setFormatter(logging.Formatter("%(asctime)s %(levelname)s %(message)s"))
    logger.addHandler(handler)
    logger.setLevel(logging.INFO)
    logger.propagate = False


def log_trace(event: str, status: str, job_id: str, **details) -> None:
    """Registra somente etapa e contagens; nunca número, cookie ou GUID."""
    fields = " ".join(f"{key}={value}" for key, value in details.items())
    severity = logging.ERROR if event.endswith("failed") else (
        logging.WARNING if event.endswith("retry") else logging.INFO
    )
    logger.log(severity, "job=%s status=%s event=%s %s", job_id, status, event, fields)


def fetch_trace(status: str, job_id: str = "manual") -> bytes:
    """Usa o login existente e os controles reais do ASP.NET/DevExpress."""
    if status not in ("19", "21"):
        raise ValueError("Status IPBX inválido")
    load_dotenv(ROOT / ".env")
    user, password = os.getenv("ORBE_LOGIN"), os.getenv("ORBE_PASSWORD")
    if not user or not password:
        raise RuntimeError("Credenciais ORBE_LOGIN/ORBE_PASSWORD ausentes")
    project = os.getenv("TOI_TRACE_PROJECT", "868")
    menu = os.getenv("TOI_TRACE_MENU", "7952")
    campaign = os.getenv("TOI_TRACE_CAMPAIGN", "1474")
    for attempt in range(3):
        portal = PortalIPBXCanaisTravados(
            usuario=user, senha=password,
            config_instancias={"trace": {"projeto": project, "menu": menu}},
        )
        try:
            log_trace("portal_login_start", status, job_id, attempt=attempt + 1)
            portal.autenticar()
            log_trace("portal_project_start", status, job_id, attempt=attempt + 1)
            _, home = portal._selecionar_projeto("trace")
            break
        except (requests.ConnectionError, requests.Timeout) as exc:
            log_trace("portal_connection_retry", status, job_id,
                      attempt=attempt + 1, error=type(exc).__name__)
            if attempt == 2:
                raise RuntimeError("Conexão com o Portal Orbe indisponível após três tentativas") from None
            time.sleep((2, 5)[attempt])
    log_trace("portal_project_ready", status, job_id)
    anchor = next(
        (a for a in BeautifulSoup(home.text, "html.parser").find_all("a")
         if " ".join(a.stripped_strings) == "Trace de Chamadas"), None
    )
    if anchor is None:
        raise RuntimeError("Menu Trace de Chamadas indisponível para a conta configurada")
    link = re.search(
        r'''openReport\s*\(\s*['"][^'"]*['"]\s*,\s*['"]([^'"]+)''',
        anchor.get("href", ""), re.I,
    )
    if link is None:
        raise RuntimeError("Link do Trace de Chamadas não reconhecido")
    trace_url = urljoin(home.url, unescape(link.group(1)))
    log_trace("trace_form_start", status, job_id)
    phase = "browser_start"
    try:
        with sync_playwright() as playwright:
            browser = playwright.chromium.launch(headless=True)
            try:
                phase = "trace_navigation"
                context = browser.new_context(accept_downloads=True)
                context.add_cookies([
                    {"name": c.name, "value": c.value, "domain": c.domain,
                     "path": c.path or "/", "secure": c.secure, "httpOnly": False}
                    for c in portal.session.cookies
                ])
                page = context.new_page()
                page.goto(trace_url, wait_until="domcontentloaded", timeout=30_000)
                if "DynamicQuery.aspx" not in page.url:
                    raise RuntimeError("O Portal não abriu o formulário do Trace")
                phase = "trace_filters"
                page.locator('[name="ctl09$comboBox"]').select_option(campaign)
                page.wait_for_function(
                    'document.querySelectorAll("[name=\'ctl36$Lst_ddlCodeStatusIPBX\'] option").length > 0',
                    timeout=30_000,
                )
                page.locator('[name="ctl15$textBox"]').fill("500")
                page.locator('[name="ctl36$Chk_ddlCodeStatusIPBX"]').check()
                page.locator('[name="ctl36$Lst_ddlCodeStatusIPBX"]').select_option(status)
                page.locator("#cmbIdFormatTypeTop_B-1").click()
                page.locator("#cmbIdFormatTypeTop_DDD_L_LBI2T").click()
                if page.locator('[name="cmbIdFormatTypeTop"]').input_value() != "XLS":
                    raise RuntimeError("O Portal não selecionou o formato XLS")
                phase = "trace_download"
                with page.expect_download(timeout=90_000) as transfer:
                    page.locator('[name="btnQueryTop"]').click()
                output = Path(transfer.value.path()).read_bytes()
                if not output.startswith(OLE_SIGNATURE):
                    raise RuntimeError("O Portal não entregou uma planilha XLS válida")
                log_trace("trace_export_ready", status, job_id, bytes=len(output))
                return output
            finally:
                browser.close()
    except Exception as exc:
        log_trace("trace_export_failed", status, job_id, phase=phase, error=type(exc).__name__)
        raise


def _carrier_name(label: str) -> str:
    label = " ".join(label.split())
    upper = label.upper()
    if "TELEF" in upper or "BETA" in upper:
        return "Beta"
    if "GAMA" in upper:
        return "Gama"
    if "ALFA" in upper:
        return "ALFA"
    if "DELTA " in upper or upper.startswith("DELTA"):
        return "Delta"
    if not label or "�" in label:
        raise RuntimeError("O site não devolveu uma operadora legível")
    return label


def lookup_carrier(session: requests.Session, phone: str, *, job_id: str = "manual",
                   status: str = "?", index: int = 0) -> str:
    for attempt in range(3):
        try:
            response = session.post(
                CONSULTA_URL, data={"tipo": "consulta", "numero": phone},
                headers={"Connection": "close"}, timeout=25,
            )
            response.raise_for_status()
            break
        except (requests.ConnectionError, requests.Timeout) as exc:
            log_trace("carrier_connection_retry", status, job_id,
                      index=index, attempt=attempt + 1, error=type(exc).__name__)
            if attempt == 2:
                raise RuntimeError("Consulta Operadora interrompeu a conexão após três tentativas") from None
            session.close()
            time.sleep((15, 30)[attempt])
    soup = BeautifulSoup(response.content, "html.parser")
    label = next((x for x in soup.select("span")
                  if x.get_text(" ", strip=True).strip().lower() == "operadora:"), None)
    value = label.find_next_sibling("span") if label else None
    if value is None:
        raise RuntimeError("Consulta Operadora não devolveu a operadora atual")
    return _carrier_name(value.get_text(" ", strip=True))


def build_current_report(status: str, *, trace_bytes: bytes | None = None,
                         delay=None, progress=None, job_id: str = "manual") -> tuple[bytes, dict]:
    """Seleciona até 100 telefones distintos e consulta um a um com 3–5 s de intervalo."""
    log_trace("job_start", status, job_id, source="upload" if trace_bytes is not None else "portal")
    contents = trace_bytes if trace_bytes is not None else fetch_trace(status, job_id)
    selection = analyze(contents, status, classifier=lambda _: ("PENDENTE", "consulta_web"))
    log_trace("sample_ready", status, job_id, rows=selection["matching_rows"],
              distinct=selection["valid_distinct"], selected=selection["selected"])
    pause = delay or (lambda: time.sleep(random.uniform(3, 5)))
    session = requests.Session()
    records = selection["records"]
    for index, record in enumerate(records):
        if index:
            pause()
        record["carrier"] = lookup_carrier(session, record["phone"],
                                             job_id=job_id, status=status, index=index + 1)
        if (index + 1) % 10 == 0:
            log_trace("carrier_progress", status, job_id, completed=index + 1,
                      selected=len(records))
        if progress:
            progress(index + 1, len(records))

    book = Workbook()
    sheet = book.active
    sheet.title = f"IPBX {status}"
    sheet.append(["DDD", "Telefone", "Operadora atual"])
    for record in records:
        sheet.append([record["ddd"], record["phone"][2:], record["carrier"]])
        sheet.cell(sheet.max_row, 1).number_format = "@"
        sheet.cell(sheet.max_row, 2).number_format = "@"
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="0B2636")
        cell.font = Font(color="FFFFFF", bold=True)
    sheet.column_dimensions["A"].width = 10
    sheet.column_dimensions["B"].width = 20
    sheet.column_dimensions["C"].width = 30
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.sheet_view.showGridLines = False
    output = BytesIO()
    book.save(output)
    log_trace("job_complete", status, job_id, selected=len(records), bytes=output.tell())
    return output.getvalue(), selection
